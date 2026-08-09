"""Supabase 사전계산 데이터를 사용하는 단일 현장 예찰 보고서 생성기."""
from __future__ import annotations

import io
import math
import os
import random
import re
import zipfile
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Mm
from docx.text.paragraph import Paragraph
from dotenv import load_dotenv
from PIL import Image, ImageChops, ImageDraw, ImageFont, ImageStat

# pyproj / shapely 대신 순수 파이썬 구현을 쓴다.
# 이 파일이 두 패키지의 유일한 사용처였고, 쓰던 기능은 EPSG:5186->4326 변환과
# 좌표 그릇(.x/.y/.centroid)뿐이었다. numpy까지 딸려와 배포 번들이
# Vercel 225MB 한도를 넘겼기 때문에 app/core/coords.py 로 대체했다.
# (한국 전역 2,556개 지점 대조에서 pyproj 대비 최대 오차 0.059mm)
from app.core.coords import Point, Transformer

from app.services.prediction_report_generator import (
    BACKEND_ROOT,
    DEFAULT_ZOOM,
    MAP_HEIGHT,
    MAP_WIDTH,
    PROJECT_ROOT,
    RISK_COLORS,
    ReportRecord,
    ReportSourceData,
    _create_supabase_client,
    calculate_metrics,
    geometry_to_image_rings,
    load_report_source_data,
    local_risk_grade,
    lonlat_to_world_pixel,
    request_vworld_map,
)

GRID_SIZE_M = 500
TEMPLATE_PATH = (
    BACKEND_ROOT
    / "data"
    / "report_templates"
    / "[양식]소나무재선충병 현장 예찰 보고서_빈양식.docx"
)
DEFAULT_OUTPUT_ROOT = (
    BACKEND_ROOT / "data" / "generated_drafts" / "field_survey_template_test"
)


@dataclass
class TerrainCell:
    grid_id: int
    geometry: Point
    pine_ratio: float = 0.0
    elev_mean: float = 0.0
    slope_mean: float = 0.0


def apply_field_candidate_metrics(
    metrics: dict[str, Any],
    candidate: dict[str, Any] | None,
) -> dict[str, Any]:
    """원본 field_survey manifest 연결 방식처럼 표시 점수만 덮어쓴다."""
    linked = dict(metrics)
    if not candidate:
        return linked
    for key in ("risk_score", "risk_grade", "priority_score", "priority_grade"):
        value = candidate.get(key)
        if value is not None:
            linked[key] = (
                float(value)
                if key in {"risk_score", "priority_score"}
                else str(value)
            )
    return linked


def load_field_survey_source_data(
    center_grid_id: int,
    year: int,
    *,
    client: Any | None = None,
) -> ReportSourceData:
    client = client or _create_supabase_client()
    return load_report_source_data(
        center_grid_id,
        year,
        client=client,
    )


@dataclass
class FieldSurveyData:
    survey_datetime: str
    weather: str
    temperature_c: float
    wind_speed_ms: float
    organization_team: str
    surveyors: str
    discovery_route: str
    address: str
    forest_compartment: str
    latitude: float
    longitude: float
    ai_result: str
    overall_judgment: str
    species: str
    total_trees: int
    detail_classification: str
    tree_height_m: float
    dbh_cm: float
    discoloration_stage: str
    vector_trace: str
    bark_wood_observation: str
    investigator_opinion: str
    sample_description: str
    qr_code: str
    system_link_result: str
    followup_date: str
    inspection_agency_status: str
    followup_plan: str
    route_type: str
    sample_count: int
    suspicious_count: int

def sanitize_filename(value: str) -> str:
    return re.sub(r'[\\/:*?"<>|]+', "_", value).strip()

def find_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    candidates = [
        "/usr/share/fonts/truetype/nanum/NanumGothicBold.ttf" if bold else "/usr/share/fonts/truetype/nanum/NanumGothic.ttf",
        "/usr/share/fonts/truetype/nanum/NanumSquareRoundB.ttf" if bold else "/usr/share/fonts/truetype/nanum/NanumSquareRoundR.ttf",
        "/usr/share/fonts/truetype/unfonts-core/UnDotumBold.ttf" if bold else "/usr/share/fonts/truetype/unfonts-core/UnDotum.ttf",
    ]
    for path in candidates:
        if Path(path).exists():
            return ImageFont.truetype(path, size=size)
    return ImageFont.load_default()

def deterministic_rng(record: ReportRecord) -> random.Random:
    return random.Random(record.center_grid_id * 10000 + record.year)

def add_days_text(year: int, month: int, day: int, offset: int) -> str:
    # 12월 날짜만 사용하며 28일을 넘지 않도록 순환
    adjusted = ((day - 1 + offset) % 28) + 1
    return f"{year}. {month:02d}. {adjusted:02d}."

def create_field_survey_data(
    record: ReportRecord,
    metrics: dict[str, Any],
    center_cell: TerrainCell,
) -> FieldSurveyData:
    rng = deterministic_rng(record)
    transformer = Transformer.from_crs("EPSG:5186", "EPSG:4326", always_xy=True)
    center = center_cell.geometry.centroid
    lon, lat = transformer.transform(center.x, center.y)

    risk = float(metrics["risk_score"])
    pressure = float(metrics["infection_pressure"])
    pine = float(metrics["pine_mean"])

    # 위험도별 관찰 강도
    if risk >= 85:
        suspicious = rng.randint(5, 9)
        route_type = "현장 정밀예찰 및 드론 보조 촬영"
        discoloration = rng.choice([
            "수관 상부 황화와 부분 갈변이 함께 관찰되는 중기 단계",
            "가지 끝부터 갈변이 확산된 중기~후기 단계",
        ])
        vector_trace = rng.choice([
            "일부 수간에서 소형 천공 흔적이 관찰되어 추가 확인 필요",
            "수피 표면에 천공 의심 흔적이 확인되어 정밀 관찰 실시",
        ])
        sample_count = rng.randint(3, 5)
    elif risk >= 70:
        suspicious = rng.randint(2, 6)
        route_type = "드론 선행예찰 후 현장 확인"
        discoloration = rng.choice([
            "수관 일부 황화가 확인되는 초기~중기 단계",
            "부분 갈변 및 잎 처짐이 관찰되는 초기 단계",
        ])
        vector_trace = rng.choice([
            "뚜렷한 탈출공은 미확인되었으나 일부 천공 의심 흔적 관찰",
            "수피 틈과 가지 분지부를 중심으로 추가 확인 필요",
        ])
        sample_count = rng.randint(2, 4)
    elif risk >= 55:
        suspicious = rng.randint(1, 3)
        route_type = "드론 우선예찰 및 표본 현장점검"
        discoloration = rng.choice([
            "경미한 잎 변색이 확인되는 초기 단계",
            "부분적인 황화가 있으나 계절성 변화 여부 추가 확인 필요",
        ])
        vector_trace = "뚜렷한 매개충 흔적은 확인되지 않음"
        sample_count = rng.randint(1, 2)
    else:
        suspicious = rng.randint(0, 2)
        route_type = "정기 순찰 및 드론 모니터링"
        discoloration = "대부분 정상이며 일부 개체에서 경미한 변색 관찰"
        vector_trace = "매개충 흔적 미확인"
        sample_count = 1 if suspicious > 0 else 0

    normal_count = rng.randint(8, 18)
    total = suspicious + normal_count
    dead_count = max(0, suspicious - rng.randint(0, 2))
    yellow_count = suspicious - dead_count

    species = rng.choice(["소나무", "곰솔", "잣나무", "리기다소나무"])
    height = round(rng.uniform(8.0, 17.5), 1)
    dbh = round(rng.uniform(18.0, 37.0), 1)
    temperature = round(rng.uniform(4.0, 17.0), 1)
    wind = round(rng.uniform(0.8, 4.8), 1)
    weather = rng.choice(["맑음", "구름 조금", "흐림", "맑은 후 구름 많음"])

    team_number = (record.report_no % 4) + 1
    team_names = [
        ("산림보호 예찰 1조", "김도윤, 이서진"),
        ("산림보호 예찰 2조", "박지훈, 최유나"),
        ("산림병해충 대응 1조", "정민수, 한지우"),
        ("현장예찰 지원 2조", "강서준, 윤가은"),
    ]
    organization, surveyors = team_names[record.report_no % len(team_names)]

    route = (
        f"AI 우선 예찰 검토지역 지정 후 {route_type} 방식으로 중심 격자와 주변 8개 격자를 순차 확인"
    )
    compartment = f"{(record.center_grid_id % 90) + 10}임반-{(record.center_grid_id % 9) + 1}소반"

    if suspicious > 0:
        bark_obs = rng.choice([
            "의심 개체의 수피 일부 건조와 목질부 수분 저하가 관찰됨",
            "수피 박리부 주변에서 갈변이 확인되어 시료 채취 실시",
            "가지 절단면 일부에서 변색이 관찰되어 목편 시료 확보",
        ])
        opinion = (
            f"중심 격자를 포함한 3×3 권역에서 변색 또는 고사 의심 개체 {suspicious}본을 확인했습니다. "
            f"현장만으로 감염 여부를 판단하기 어려워 시료를 채취했으며, 인접 격자까지 추가 모니터링이 필요합니다."
        )
    else:
        bark_obs = "수피 및 목질부에서 특이사항이 확인되지 않음"
        opinion = (
            "현재 조사 범위에서는 뚜렷한 감염 의심 징후가 확인되지 않았습니다. "
            "다만 과거 발생 이력이 있는 권역이므로 정기 모니터링을 유지하겠습니다."
        )

    qr_code = f"FS-{record.year}-{record.center_grid_id}-{record.report_no:02d}"
    if sample_count > 0:
        sample_desc = f"목편 {sample_count}점, 가지 시료 {max(1, sample_count - 1)}점"
        system_result = f"현장 예찰 시스템에 사진·좌표·시료정보 등록 완료, QR코드 {qr_code} 연동"
        agency_status = "관할 산림환경연구기관 검경 의뢰 접수, 결과 대기"
        followup_plan = (
            "검경 결과 확인 전까지 대상 3×3 권역을 우선 예찰 검토지역으로 유지하고, "
            "의심 개체 주변 반경을 재확인한다. 검경 결과에 따라 방제 검토와 인접 권역 확대조사를 연계한다."
        )
    else:
        sample_desc = "현장 관찰 결과 시료 미채취"
        system_result = f"현장 사진·좌표·조사결과 등록 완료, 조사기록 {qr_code} 생성"
        agency_status = "검경 의뢰 없음, 정기 예찰 대상으로 관리"
        followup_plan = (
            "2주 이내 동일 권역을 재확인하고 변색 진행 여부를 비교한다. "
            "새로운 의심 징후 확인 시 즉시 시료 채취와 검경 의뢰를 시행한다."
        )

    survey_day = 8 + (record.report_no % 18)
    survey_datetime = f"{record.year}. 12. {survey_day:02d}. 09:30~14:30"
    followup_date = add_days_text(record.year, 12, survey_day, 7)

    detail = f"정상 관찰 {normal_count}본, 변색 의심 {yellow_count}본, 고사 의심 {dead_count}본"
    ai_result = (
        f"종합 위험도 {risk:.1f}점({metrics['risk_grade']}), 감염압력 {pressure:.1f}점, "
        f"예찰 우선순위 {metrics['priority_score']:.1f}점({metrics['priority_grade']})"
    )
    judgment = (
        f"과거 발생 이력과 3×3 권역 내 위험 신호를 고려할 때 현장 확인이 필요함. "
        f"권역 평균 소나무류 비율은 {pine:.1f}%이며, 중심 격자와 인접 격자의 연속 예찰을 권고함."
    )

    return FieldSurveyData(
        survey_datetime=survey_datetime,
        weather=weather,
        temperature_c=temperature,
        wind_speed_ms=wind,
        organization_team=organization,
        surveyors=surveyors,
        discovery_route=route,
        address=f"{record.sido_name} {record.sigungu_name} 산림 일원",
        forest_compartment=compartment,
        latitude=lat,
        longitude=lon,
        ai_result=ai_result,
        overall_judgment=judgment,
        species=species,
        total_trees=total,
        detail_classification=detail,
        tree_height_m=height,
        dbh_cm=dbh,
        discoloration_stage=discoloration,
        vector_trace=vector_trace,
        bark_wood_observation=bark_obs,
        investigator_opinion=opinion,
        sample_description=sample_desc,
        qr_code=qr_code,
        system_link_result=system_result,
        followup_date=followup_date,
        inspection_agency_status=agency_status,
        followup_plan=followup_plan,
        route_type=route_type,
        sample_count=sample_count,
        suspicious_count=suspicious,
    )

def draw_text_fit(
    draw: ImageDraw.ImageDraw,
    box: tuple[int, int, int, int],
    text: str,
    *,
    max_size: int = 22,
    min_size: int = 11,
    bold: bool = False,
    align: str = "center",
    line_spacing: int = 3,
    fill: tuple[int, int, int, int] | tuple[int, int, int] = (20, 20, 20),
) -> None:
    """지정 셀 안에서 자동 줄바꿈/축소해 텍스트를 그린다."""
    x1, y1, x2, y2 = box
    width = max(1, x2 - x1 - 8)
    height = max(1, y2 - y1 - 6)
    raw = str(text).strip()
    if not raw:
        return

    def wrap_text(font: ImageFont.ImageFont) -> list[str]:
        lines: list[str] = []
        for source_line in raw.splitlines() or [""]:
            if not source_line:
                lines.append("")
                continue
            current = ""
            for ch in source_line:
                trial = current + ch
                bbox = draw.textbbox((0, 0), trial, font=font)
                if bbox[2] - bbox[0] <= width or not current:
                    current = trial
                else:
                    lines.append(current)
                    current = ch
            if current:
                lines.append(current)
        return lines or [""]

    chosen_font = find_font(min_size, bold=bold)
    chosen_lines = [raw]
    for size in range(max_size, min_size - 1, -1):
        font = find_font(size, bold=bold)
        lines = wrap_text(font)
        line_heights = []
        for line in lines:
            bbox = draw.textbbox((0, 0), line or "가", font=font)
            line_heights.append(max(1, bbox[3] - bbox[1]))
        total_h = sum(line_heights) + line_spacing * max(0, len(lines) - 1)
        if total_h <= height:
            chosen_font = font
            chosen_lines = lines
            break

    line_metrics = []
    for line in chosen_lines:
        bbox = draw.textbbox((0, 0), line or "가", font=chosen_font)
        line_metrics.append((bbox[2] - bbox[0], max(1, bbox[3] - bbox[1])))
    total_h = sum(h for _, h in line_metrics) + line_spacing * max(0, len(chosen_lines) - 1)
    y = y1 + max(0, (height - total_h) / 2) + 2

    for line, (line_w, line_h) in zip(chosen_lines, line_metrics):
        if align == "left":
            x = x1 + 5
        elif align == "right":
            x = x2 - line_w - 5
        else:
            x = x1 + max(0, ((x2 - x1) - line_w) / 2)
        draw.text((x, y), line, font=chosen_font, fill=fill)
        y += line_h + line_spacing

def extract_template_media(template_path: Path, media_name: str) -> Image.Image:
    """DOCX 템플릿의 word/media 이미지를 읽는다."""
    internal = f"word/media/{media_name}"
    with zipfile.ZipFile(template_path, "r") as archive:
        if internal not in archive.namelist():
            raise RuntimeError(f"템플릿에서 별지 이미지를 찾지 못했습니다: {internal}")
        raw = archive.read(internal)
    return Image.open(io.BytesIO(raw)).convert("RGB")

def build_air_survey_plan_appendix(
    template_path: Path,
    output_path: Path,
    record: ReportRecord,
    survey: FieldSurveyData,
    metrics: dict[str, Any],
) -> None:
    """7페이지: 유인항공예찰 계획 별지 작성."""
    image = extract_template_media(template_path, "image1.png")
    if image.size != (826, 1264):
        raise RuntimeError(
            f"유인항공예찰 계획 이미지 크기가 예상과 다릅니다: {image.size}, 예상=(826, 1264)"
        )
    draw = ImageDraw.Draw(image)

    # 첫 번째 데이터 행 좌표(원본 826×1264 기준)
    boxes = {
        "sigungu": (24, 239, 130, 286),
        "region": (130, 239, 210, 286),
        "area": (210, 239, 295, 286),
        "date": (295, 239, 374, 286),
        "landing": (374, 239, 470, 286),
        "org": (470, 239, 553, 286),
        "rank": (553, 239, 637, 286),
        "name": (637, 239, 722, 286),
        "helicopter": (722, 239, 814, 286),
        "attachment": (24, 1115, 814, 1193),
    }

    # 원본 양식에 들어 있던 파란 안내문 및 기본 문구를 지운 뒤 셀 내부만 다시 작성
    for box in boxes.values():
        x1, y1, x2, y2 = box
        draw.rectangle((x1 + 2, y1 + 2, x2 - 2, y2 - 2), fill=(255, 255, 255))

    survey_date = survey.survey_datetime.split(". 09:")[0].strip()
    area_ha = 9 * (GRID_SIZE_M * GRID_SIZE_M) / 10_000
    primary_name = survey.surveyors.split(",")[0].strip()
    org_name = "산림보호과"
    risk_note = (
        f"중심 격자 {record.center_grid_id}\n"
        f"주변 8개 포함"
    )

    draw_text_fit(draw, boxes["sigungu"], record.sigungu_name, max_size=18, min_size=12, bold=True)
    draw_text_fit(draw, boxes["region"], risk_note, max_size=15, min_size=10)
    draw_text_fit(draw, boxes["area"], f"{area_ha:.0f}", max_size=18, min_size=12)
    draw_text_fit(draw, boxes["date"], survey_date, max_size=15, min_size=10)
    draw_text_fit(draw, boxes["landing"], f"{record.sigungu_name}\n임시착륙장", max_size=14, min_size=10)
    draw_text_fit(draw, boxes["org"], org_name, max_size=14, min_size=10)
    draw_text_fit(draw, boxes["rank"], "주무관", max_size=16, min_size=11)
    draw_text_fit(draw, boxes["name"], primary_name, max_size=16, min_size=11)
    draw_text_fit(draw, boxes["helicopter"], "산림청\n중형헬기", max_size=14, min_size=10)
    draw_text_fit(
        draw,
        boxes["attachment"],
        (
            f"중심 격자 {record.center_grid_id} 포함 3×3 예찰권역, "
            f"계획면적 {area_ha:.0f}ha. 본문 현장 예찰 조사도 및 이동경로 참조."
        ),
        max_size=17,
        min_size=12,
        align="left",
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    image.save(output_path, format="PNG")

def build_air_survey_result_appendix(
    template_path: Path,
    output_path: Path,
    record: ReportRecord,
    survey: FieldSurveyData,
    metrics: dict[str, Any],
) -> None:
    """8페이지: 유인항공예찰 조사결과 별지 작성."""
    image = extract_template_media(template_path, "image2.png")
    if image.size != (832, 1228):
        raise RuntimeError(
            f"유인항공예찰 조사결과 이미지 크기가 예상과 다릅니다: {image.size}, 예상=(832, 1228)"
        )
    draw = ImageDraw.Draw(image)

    # 상단 기본정보
    # 상단 빈칸에 남아 있는 원본 예시문을 먼저 제거
    draw.rectangle((145, 143, 810, 190), fill=(255, 255, 255))
    draw.rectangle((150, 198, 810, 245), fill=(255, 255, 255))
    draw.rectangle((300, 250, 810, 296), fill=(255, 255, 255))

    draw_text_fit(
        draw,
        (145, 145, 785, 188),
        f"{record.sido_name} {record.sigungu_name} 산림보호 담당부서",
        max_size=21,
        min_size=14,
        align="left",
        bold=True,
    )
    survey_date = survey.survey_datetime.split(" 09:")[0].strip()
    draw_text_fit(
        draw,
        (150, 201, 790, 242),
        f"{survey_date} ~ {survey_date} (1일)",
        max_size=20,
        min_size=13,
        align="left",
    )
    draw_text_fit(
        draw,
        (300, 253, 790, 293),
        "1대(산림청 중형헬기)",
        max_size=20,
        min_size=13,
        align="left",
    )

    # 첫 번째 결과 행
    boxes = {
        "sigungu": (20, 507, 121, 584),
        "region": (121, 507, 296, 584),
        "area": (296, 507, 402, 584),
        "total": (402, 507, 495, 584),
        "pine": (495, 507, 603, 584),
        "oak": (603, 507, 710, 584),
        "note": (710, 507, 809, 584),
    }
    area_ha = 9 * (GRID_SIZE_M * GRID_SIZE_M) / 10_000
    detected = int(survey.suspicious_count)
    draw_text_fit(draw, boxes["sigungu"], record.sigungu_name, max_size=18, min_size=12, bold=True)
    draw_text_fit(
        draw,
        boxes["region"],
        f"격자 {record.center_grid_id}\n3×3 예찰권역",
        max_size=16,
        min_size=10,
    )
    draw_text_fit(draw, boxes["area"], f"{area_ha:.0f}", max_size=18, min_size=12)
    draw_text_fit(draw, boxes["total"], str(detected), max_size=20, min_size=13, bold=True)
    draw_text_fit(draw, boxes["pine"], str(detected), max_size=20, min_size=13)
    draw_text_fit(draw, boxes["oak"], "0", max_size=20, min_size=13)
    note = (
        f"시료 {survey.sample_count}점\n현장 확인 필요"
        if survey.sample_count > 0
        else "이상징후 미미\n정기 관찰"
    )
    draw_text_fit(draw, boxes["note"], note, max_size=14, min_size=10)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    image.save(output_path, format="PNG")

def replace_docx_media(docx_path: Path, replacements: dict[str, Path]) -> None:
    """DOCX ZIP 안의 지정 media 파일을 새 PNG로 안전하게 교체한다."""
    temp_path = docx_path.with_suffix(".media-patched.docx")
    internal_replacements = {
        f"word/media/{name}": path for name, path in replacements.items()
    }

    with zipfile.ZipFile(docx_path, "r") as src, zipfile.ZipFile(
        temp_path, "w", zipfile.ZIP_DEFLATED
    ) as dst:
        names = set(src.namelist())
        missing = set(internal_replacements) - names
        if missing:
            raise RuntimeError(
                "DOCX에서 교체 대상 이미지를 찾지 못했습니다: " + ", ".join(sorted(missing))
            )
        for item in src.infolist():
            if item.filename in internal_replacements:
                dst.writestr(item, internal_replacements[item.filename].read_bytes())
            else:
                dst.writestr(item, src.read(item.filename))

    if not temp_path.exists() or temp_path.stat().st_size == 0:
        raise RuntimeError("별지 이미지 교체 후 DOCX 생성에 실패했습니다.")
    temp_path.replace(docx_path)

def replace_in_paragraph(paragraph: Any, replacements: dict[str, str]) -> None:
    full = "".join(run.text for run in paragraph.runs)
    if not full:
        return
    new = full
    for old, value in replacements.items():
        new = new.replace(old, str(value))
    if new == full:
        return
    if paragraph.runs:
        paragraph.runs[0].text = new
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.text = new


def replace_everywhere(doc: Document, replacements: dict[str, str]) -> None:
    for paragraph in doc.paragraphs:
        replace_in_paragraph(paragraph, replacements)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    replace_in_paragraph(paragraph, replacements)
    for section in doc.sections:
        for paragraph in section.header.paragraphs:
            replace_in_paragraph(paragraph, replacements)
        for paragraph in section.footer.paragraphs:
            replace_in_paragraph(paragraph, replacements)


def set_paragraph(doc: Document, prefix: str, text: str) -> None:
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            if paragraph.runs:
                paragraph.runs[0].text = text
                for run in paragraph.runs[1:]:
                    run.text = ""
            else:
                paragraph.text = text
            return
    raise RuntimeError(f"템플릿에서 문단을 찾지 못했습니다: {prefix}")


def insert_paragraph_after(paragraph: Paragraph) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    return Paragraph(new_p, paragraph._parent)

def insert_map_after_prefix(doc: Document, prefix: str, map_path: Path) -> None:
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            map_paragraph = insert_paragraph_after(paragraph)
            map_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = map_paragraph.add_run()
            run.add_picture(str(map_path), width=Mm(165))
            caption = insert_paragraph_after(map_paragraph)
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            caption.add_run("[현장 예찰 조사도 - 중심 격자 및 주변 3×3 권역]")
            return
    raise RuntimeError(f"지도 삽입 기준 문단을 찾지 못했습니다: {prefix}")

def create_docx(
    template_path: Path,
    output_path: Path,
    map_path: Path,
    plan_appendix_path: Path,
    result_appendix_path: Path,
    record: ReportRecord,
    cells: list[TerrainCell],
    metrics: dict[str, Any],
    survey: FieldSurveyData,
) -> None:
    doc = Document(template_path)
    region = f"{record.sido_name} {record.sigungu_name}"

    replacements = {
        "[작성일]": f"{record.year}. 12. {10 + (record.report_no % 18):02d}.",
        "-지역, 기간-": f"-{region}, {record.year}년-",
        "[일시]": survey.survey_datetime,
        "[날씨]": survey.weather,
        "기온 [ ]℃": f"기온 {survey.temperature_c:.1f}℃",
        "풍속 [ ]m/s": f"풍속 {survey.wind_speed_ms:.1f}m/s",
        "[소속·조]": survey.organization_team,
        "[성명]": survey.surveyors,
        "[입력]": survey.discovery_route,
        "[주소]": survey.address,
        "[정보]": survey.forest_compartment,
        "[위도]": f"{survey.latitude:.6f}",
        "[경도]": f"{survey.longitude:.6f}",
        "[점수·등급 또는 판단 내용]": survey.ai_result,
        "[현장 확인 필요 여부 및 판단 근거 입력]": survey.overall_judgment,
        "[수종]": survey.species,
        "[수량]": str(survey.total_trees),
        "[세부 분류]": survey.detail_classification,
        "수고 약 [ ]m": f"수고 약 {survey.tree_height_m:.1f}m",
        "흉고직경(DBH) [ ]cm": f"흉고직경(DBH) {survey.dbh_cm:.1f}cm",
        "[관찰 내용 및 단계]": survey.discoloration_stage,
        "[관찰 내용]": survey.vector_trace,
        "[현장 조사자 의견 입력]": survey.investigator_opinion,
        "[시료 종류·수량]": survey.sample_description,
        "[번호]": survey.qr_code,
        "[처리 결과 입력]": survey.system_link_result,
        "[일자]": survey.followup_date,
        "[검경 의뢰 기관 및 진행 상태]": survey.inspection_agency_status,
        "[현장 확인·검경 결과에 따른 조치 입력]": survey.followup_plan,
    }
    replace_everywhere(doc, replacements)

    # 동일한 [관찰 내용] placeholder가 2개라 명시적으로 다시 설정
    set_paragraph(doc, "❍ (변색 단계)", f"❍ (변색 단계) {survey.discoloration_stage}")
    set_paragraph(doc, "❍ (매개충 흔적)", f"❍ (매개충 흔적) {survey.vector_trace}")
    set_paragraph(doc, "❍ (수피·목질부 관찰)", f"❍ (수피·목질부 관찰) {survey.bark_wood_observation}")

    # 기존 양식의 GPS 문단 바로 뒤에 지도 삽입
    insert_map_after_prefix(doc, "― 세부 좌표:", map_path)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)

    # 템플릿의 7~8페이지 별지 이미지(image1, image2)를 채워진 이미지로 교체
    replace_docx_media(
        output_path,
        {
            "image1.png": plan_appendix_path,
            "image2.png": result_appendix_path,
        },
    )


def verify_appendix_image_changed(
    template_path: Path,
    output_path: Path,
    media_name: str,
) -> None:
    """NumPy 없이 Pillow만으로 별지 이미지 변경 여부를 검증한다."""
    original = extract_template_media(template_path, media_name)
    filled = Image.open(output_path).convert("RGB")
    if original.size != filled.size:
        raise RuntimeError(
            f"별지 이미지 크기가 바뀌었습니다: {media_name}"
        )
    difference = ImageChops.difference(original, filled)
    mean_difference = sum(ImageStat.Stat(difference).mean) / 3.0
    if mean_difference < 0.15:
        raise RuntimeError(
            f"별지 이미지에 입력 내용이 충분히 그려지지 않았습니다: "
            f"{output_path}"
        )


def make_survey_points(
    record: ReportRecord,
    center_cell: TerrainCell,
    survey: FieldSurveyData,
) -> tuple[list[Point], list[Point], list[Point]]:
    rng = deterministic_rng(record)
    center = center_cell.geometry.centroid
    route_points = [
        Point(center.x - 600, center.y - 500),
        Point(center.x - 350, center.y + 450),
        Point(center.x + 100, center.y + 550),
        Point(center.x + 550, center.y + 250),
        Point(center.x + 500, center.y - 400),
        Point(center.x, center.y - 550),
    ]
    observation_count = max(1, min(5, survey.suspicious_count))
    observation_points = [
        Point(
            center.x + rng.uniform(-620, 620),
            center.y + rng.uniform(-620, 620),
        )
        for _ in range(observation_count)
    ]
    sample_points = observation_points[
        : min(survey.sample_count, len(observation_points))
    ]
    return route_points, observation_points, sample_points


def projected_point_to_image(
    point_5186: Point,
    center_lon: float,
    center_lat: float,
    zoom: int,
    transformer: Transformer,
) -> tuple[float, float]:
    lon, lat = transformer.transform(point_5186.x, point_5186.y)
    center_x, center_y = lonlat_to_world_pixel(
        center_lon,
        center_lat,
        zoom,
    )
    world_x, world_y = lonlat_to_world_pixel(lon, lat, zoom)
    return (
        MAP_WIDTH / 2 + (world_x - center_x),
        MAP_HEIGHT / 2 + (world_y - center_y),
    )


def _draw_geojson(
    draw: ImageDraw.ImageDraw,
    geometry: dict[str, Any],
    center_lon: float,
    center_lat: float,
    zoom: int,
    *,
    fill: tuple[int, int, int, int] | None = None,
    outline: tuple[int, int, int, int],
    width: int,
) -> list[list[tuple[float, float]]]:
    rings = geometry_to_image_rings(
        geometry,
        center_lon,
        center_lat,
        zoom,
    )
    for points in rings:
        if fill is not None:
            draw.polygon(points, fill=fill, outline=outline)
        else:
            draw.line(points, fill=outline, width=width, joint="curve")
        if fill is not None and width > 1:
            draw.line(points, fill=outline, width=width, joint="curve")
    return rings


def render_field_survey_overlay(
    background: Image.Image,
    record: ReportRecord,
    source: ReportSourceData,
    survey: FieldSurveyData,
    zoom: int,
) -> tuple[Image.Image, dict[str, Any]]:
    static = source.static
    center_4326 = static["center_point_4326"]["coordinates"]
    center_lon, center_lat = map(float, center_4326[:2])
    center_5186 = static["center_point_5186"]["coordinates"]
    center_cell = TerrainCell(
        grid_id=record.center_grid_id,
        geometry=Point(float(center_5186[0]), float(center_5186[1])),
        pine_ratio=float(static["block_pine_mean"]),
        elev_mean=float(static["block_elevation_mean"]),
        slope_mean=float(static["block_slope_mean"]),
    )
    overlay = Image.new("RGBA", background.size, (0, 0, 0, 0))
    draw = ImageDraw.Draw(overlay, "RGBA")
    annual_column = (
        f"infection_count_{record.year}"
        if 2016 <= record.year <= 2021
        else None
    )

    infection_rings: dict[str, list[list[tuple[float, float]]]] = {}
    for row in source.infection_positions:
        annual = int(row.get(annual_column, 0) or 0) if annual_column else 0
        cumulative = int(row.get("infection_count_2016_2021", 0) or 0)
        if annual <= 0 and cumulative <= 0:
            continue
        grade = local_risk_grade(annual, cumulative)
        color = RISK_COLORS[grade]
        rings = geometry_to_image_rings(
            row["geometry_4326"],
            center_lon,
            center_lat,
            zoom,
        )
        visible = any(
            -100 <= x <= MAP_WIDTH + 100
            and -100 <= y <= MAP_HEIGHT + 100
            for ring in rings
            for x, y in ring
        )
        if not visible:
            continue
        infection_rings[str(row["grid_id"])] = rings
        for points in rings:
            draw.polygon(
                points,
                fill=(color[0], color[1], color[2], 120),
                outline=(255, 255, 255, 90),
            )

    block_ids = [int(value) for value in static["block_grid_ids"]]
    block_geometries = list(static["block_cell_geometries_4326"])
    block_rings: dict[str, list[list[tuple[float, float]]]] = {}
    for grid_id, geometry in zip(block_ids, block_geometries):
        is_center = grid_id == record.center_grid_id
        block_rings[str(grid_id)] = _draw_geojson(
            draw,
            geometry,
            center_lon,
            center_lat,
            zoom,
            fill=(255, 65, 65, 75) if is_center else None,
            outline=(
                (170, 0, 0, 255)
                if is_center
                else (20, 82, 190, 240)
            ),
            width=5 if is_center else 4,
        )
    outer_rings = _draw_geojson(
        draw,
        static["block_geometry_4326"],
        center_lon,
        center_lat,
        zoom,
        outline=(0, 45, 130, 255),
        width=6,
    )

    transformer = Transformer.from_crs(
        "EPSG:5186",
        "EPSG:4326",
        always_xy=True,
    )
    route_points, observation_points, sample_points = make_survey_points(
        record,
        center_cell,
        survey,
    )
    route_pixels = [
        projected_point_to_image(
            point,
            center_lon,
            center_lat,
            zoom,
            transformer,
        )
        for point in route_points
    ]
    draw.line(
        route_pixels,
        fill=(20, 20, 20, 230),
        width=5,
        joint="curve",
    )
    for first, second in zip(route_pixels, route_pixels[1:]):
        x1, y1 = first
        x2, y2 = second
        middle_x = x1 + (x2 - x1) * 0.72
        middle_y = y1 + (y2 - y1) * 0.72
        angle = math.atan2(y2 - y1, x2 - x1)
        size = 12
        draw.polygon(
            [
                (middle_x, middle_y),
                (
                    middle_x - size * math.cos(angle - 0.5),
                    middle_y - size * math.sin(angle - 0.5),
                ),
                (
                    middle_x - size * math.cos(angle + 0.5),
                    middle_y - size * math.sin(angle + 0.5),
                ),
            ],
            fill=(20, 20, 20, 230),
        )

    observation_pixels = [
        projected_point_to_image(
            point,
            center_lon,
            center_lat,
            zoom,
            transformer,
        )
        for point in observation_points
    ]
    sample_pixels = [
        projected_point_to_image(
            point,
            center_lon,
            center_lat,
            zoom,
            transformer,
        )
        for point in sample_points
    ]
    for x, y in observation_pixels:
        draw.polygon(
            [(x, y - 13), (x - 11, y + 9), (x + 11, y + 9)],
            fill=(255, 45, 45, 245),
            outline=(120, 0, 0, 255),
        )
    for x, y in sample_pixels:
        draw.rectangle(
            (x - 8, y - 8, x + 8, y + 8),
            fill=(35, 85, 220, 255),
            outline=(255, 255, 255, 255),
            width=2,
        )
    start_x, start_y = route_pixels[0]
    draw.ellipse(
        (start_x - 10, start_y - 10, start_x + 10, start_y + 10),
        fill=(35, 190, 95, 255),
        outline=(255, 255, 255, 255),
        width=2,
    )

    title_font = find_font(30, bold=True)
    label_font = find_font(22, bold=True)
    small_font = find_font(17)
    title = f"{record.sido_name} {record.sigungu_name} / 현장 예찰 조사도"
    title_box = draw.textbbox((0, 0), title, font=title_font)
    title_width = title_box[2] - title_box[0]
    draw.rounded_rectangle(
        (
            MAP_WIDTH / 2 - title_width / 2 - 16,
            14,
            MAP_WIDTH / 2 + title_width / 2 + 16,
            60,
        ),
        radius=9,
        fill=(255, 255, 255, 230),
        outline=(150, 150, 150, 220),
        width=2,
    )
    draw.text(
        (MAP_WIDTH / 2 - title_width / 2, 20),
        title,
        font=title_font,
        fill=(20, 20, 20, 255),
    )
    center_label = f"중심 격자 {record.center_grid_id}"
    draw.rounded_rectangle(
        (
            MAP_WIDTH / 2 + 12,
            MAP_HEIGHT / 2 - 52,
            MAP_WIDTH / 2 + 250,
            MAP_HEIGHT / 2 - 14,
        ),
        radius=7,
        fill=(255, 255, 255, 235),
        outline=(160, 160, 160, 255),
        width=2,
    )
    draw.text(
        (MAP_WIDTH / 2 + 20, MAP_HEIGHT / 2 - 48),
        center_label,
        font=label_font,
        fill=(20, 20, 20, 255),
    )

    legend_x, legend_y = 20, MAP_HEIGHT - 158
    draw.rounded_rectangle(
        (legend_x, legend_y, legend_x + 560, legend_y + 135),
        radius=10,
        fill=(255, 255, 255, 238),
        outline=(165, 165, 165, 230),
        width=2,
    )
    legend = [
        ("예찰 시작점", "circle", (35, 190, 95, 255)),
        ("감염 의심목 관찰점", "triangle", (255, 45, 45, 245)),
        ("시료 채취점", "square", (35, 85, 220, 255)),
        ("현장 이동경로", "line", (20, 20, 20, 230)),
        ("중심 격자", "box_red", (255, 65, 65, 75)),
        ("주변 8개 격자", "box_blue", (20, 82, 190, 240)),
    ]
    for index, (name, shape_name, color) in enumerate(legend):
        x = legend_x + 18 + (index % 2) * 275
        y = legend_y + 17 + (index // 2) * 38
        if shape_name == "circle":
            draw.ellipse((x, y, x + 22, y + 22), fill=color)
        elif shape_name == "triangle":
            draw.polygon(
                [(x + 11, y), (x, y + 22), (x + 22, y + 22)],
                fill=color,
            )
        elif shape_name == "square":
            draw.rectangle((x, y, x + 22, y + 22), fill=color)
        elif shape_name == "line":
            draw.line((x, y + 11, x + 26, y + 11), fill=color, width=4)
        elif shape_name == "box_red":
            draw.rectangle(
                (x, y, x + 24, y + 22),
                fill=(255, 65, 65, 75),
                outline=(170, 0, 0, 255),
                width=3,
            )
        else:
            draw.rectangle(
                (x, y, x + 24, y + 22),
                fill=(255, 255, 255, 20),
                outline=(20, 82, 190, 240),
                width=3,
            )
        draw.text(
            (x + 34, y - 1),
            name,
            font=small_font,
            fill=(20, 20, 20, 255),
        )
    source_text = "배경지도: VWorld | 예찰권역: 중심 격자 포함 3×3"
    source_box = draw.textbbox((0, 0), source_text, font=small_font)
    source_width = source_box[2] - source_box[0]
    draw.rounded_rectangle(
        (
            MAP_WIDTH - source_width - 28,
            MAP_HEIGHT - 34,
            MAP_WIDTH - 10,
            MAP_HEIGHT - 6,
        ),
        radius=5,
        fill=(255, 255, 255, 225),
    )
    draw.text(
        (MAP_WIDTH - source_width - 20, MAP_HEIGHT - 32),
        source_text,
        font=small_font,
        fill=(50, 50, 50, 255),
    )
    result = Image.alpha_composite(
        background.convert("RGBA"),
        overlay,
    ).convert("RGB")
    return result, {
        "route_points_5186": [(point.x, point.y) for point in route_points],
        "observation_points_5186": [
            (point.x, point.y) for point in observation_points
        ],
        "sample_points_5186": [
            (point.x, point.y) for point in sample_points
        ],
        "route_pixels": route_pixels,
        "observation_pixels": observation_pixels,
        "sample_pixels": sample_pixels,
        "infection_polygons": infection_rings,
        "block_polygons": block_rings,
        "outer_polygons": outer_rings,
    }


def build_field_survey_map(
    *,
    output_path: Path,
    record: ReportRecord,
    source: ReportSourceData,
    survey: FieldSurveyData,
    api_key: str,
    domain: str,
    zoom: int,
    basemap: str,
) -> dict[str, Any]:
    center = source.static["center_point_4326"]["coordinates"]
    background = request_vworld_map(
        api_key=api_key,
        domain=domain,
        center_lon=float(center[0]),
        center_lat=float(center[1]),
        zoom=zoom,
        basemap=basemap,
        width=MAP_WIDTH,
        height=MAP_HEIGHT,
    )
    result, trace = render_field_survey_overlay(
        background,
        record,
        source,
        survey,
        zoom,
    )
    output_path.parent.mkdir(parents=True, exist_ok=True)
    result.save(output_path, quality=95)
    return trace


def build_field_survey_render_payload(
    *,
    record: ReportRecord,
    metrics: dict[str, Any],
    survey: FieldSurveyData,
    map_path: Path,
) -> dict[str, Any]:
    center_grid = {
        "grid_id": record.center_grid_id,
        "risk_score": metrics["risk_score"],
        "risk_grade": metrics["risk_grade"],
        "priority_score": metrics["priority_score"],
        "priority_grade": metrics["priority_grade"],
        "pine_ratio": metrics["pine_mean"] / 100.0,
        "latitude": survey.latitude,
        "longitude": survey.longitude,
    }
    data = {
        **asdict(survey),
        **center_grid,
        "center_grid_id": record.center_grid_id,
        "document_no": record.report_no,
        "year": record.year,
        "block_grid_ids": metrics["block_grid_ids"],
    }
    return {
        "report_type": "field_survey",
        "document_no": record.report_no,
        "field_report_id": (
            f"RPT-FIELD-{record.year}-{record.report_no:03d}"
        ),
        "year": record.year,
        "title": (
            f"{record.year}년 {record.sigungu_name} 현장 예찰 보고서"
        ),
        "start_date": f"{record.year}-12-{8 + record.report_no % 18:02d}",
        "end_date": f"{record.year}-12-{8 + record.report_no % 18:02d}",
        "sido_name": record.sido_name,
        "sigungu_name": record.sigungu_name,
        "center_grid_id": record.center_grid_id,
        "center_grid_ids": [str(record.center_grid_id)],
        "block_grid_ids": metrics["block_grid_ids"],
        "map_path": str(map_path),
        "field_data": data,
        "data": data,
        "data_summary": {
            "selected_grid_count": 1,
            "region_candidate_count": 1,
            "grid_ids": [str(record.center_grid_id)],
            "center_grid": center_grid,
            "neighbor_grids": [],
        },
    }


def generate_single_field_survey_report(
    *,
    center_grid_id: int,
    year: int,
    output_root: Path | None = None,
    report_no: int = 1,
    zoom: int | None = None,
    candidate_metrics: dict[str, Any] | None = None,
    client: Any | None = None,
) -> dict[str, Any]:
    if year < 2016 or year > 2100:
        raise ValueError("보고서 연도는 2016~2100 범위여야 합니다.")
    load_dotenv(BACKEND_ROOT / ".env")
    load_dotenv(PROJECT_ROOT / ".env")
    api_key = os.getenv("VWORLD_API_KEY", "").strip()
    domain = os.getenv("VWORLD_API_DOMAIN", "").strip()
    basemap = os.getenv("VWORLD_BASEMAP", "GRAPHIC").strip() or "GRAPHIC"
    if not api_key or not domain:
        raise RuntimeError("VWORLD_API_KEY와 VWORLD_API_DOMAIN이 필요합니다.")

    source = load_field_survey_source_data(
        int(center_grid_id),
        int(year),
        client=client,
    )
    historical = 2016 <= year <= 2021
    record = ReportRecord(
        report_no=int(report_no),
        year=int(year),
        center_grid_id=int(center_grid_id),
        annual_count=(
            int(source.stats["center_annual_count"]) if historical else 0
        ),
        cumulative_count=int(source.stats["center_cumulative_count"]),
        sido_name=str(source.static["sido_name"]),
        sigungu_name=str(source.static["sigungu_name"]),
    )
    metrics = apply_field_candidate_metrics(
        calculate_metrics(record, source),
        candidate_metrics,
    )
    center_5186 = source.static["center_point_5186"]["coordinates"]
    center_cell = TerrainCell(
        grid_id=record.center_grid_id,
        geometry=Point(float(center_5186[0]), float(center_5186[1])),
        pine_ratio=metrics["pine_mean"],
        elev_mean=metrics["elevation_mean"],
        slope_mean=metrics["slope_mean"],
    )
    survey = create_field_survey_data(record, metrics, center_cell)

    root = (
        output_root.resolve()
        if output_root is not None
        else DEFAULT_OUTPUT_ROOT.resolve()
    )
    directories = {
        name: root / name
        for name in ("docx", "pdf", "maps", "appendices")
    }
    for directory in [root, *directories.values()]:
        directory.mkdir(parents=True, exist_ok=True)
    base_name = sanitize_filename(
        f"{record.report_no:02d}_{record.year}_"
        f"소나무재선충병_현장예찰보고서_"
        f"{record.sido_name}_{record.sigungu_name}_"
        f"격자{record.center_grid_id}"
    )
    map_path = directories["maps"] / f"{base_name}.png"
    docx_path = directories["docx"] / f"{base_name}.docx"
    pdf_path = directories["pdf"] / f"{base_name}.pdf"
    plan_appendix = (
        directories["appendices"] / f"{base_name}_07_유인항공예찰계획.png"
    )
    result_appendix = (
        directories["appendices"] / f"{base_name}_08_유인항공예찰조사결과.png"
    )
    selected_zoom = zoom or int(
        os.getenv("VWORLD_ZOOM", str(DEFAULT_ZOOM))
    )
    build_field_survey_map(
        output_path=map_path,
        record=record,
        source=source,
        survey=survey,
        api_key=api_key,
        domain=domain,
        zoom=selected_zoom,
        basemap=basemap,
    )
    build_air_survey_plan_appendix(
        TEMPLATE_PATH,
        plan_appendix,
        record,
        survey,
        metrics,
    )
    build_air_survey_result_appendix(
        TEMPLATE_PATH,
        result_appendix,
        record,
        survey,
        metrics,
    )
    verify_appendix_image_changed(
        TEMPLATE_PATH,
        plan_appendix,
        "image1.png",
    )
    verify_appendix_image_changed(
        TEMPLATE_PATH,
        result_appendix,
        "image2.png",
    )
    create_docx(
        TEMPLATE_PATH,
        docx_path,
        map_path,
        plan_appendix,
        result_appendix,
        record,
        [center_cell],
        metrics,
        survey,
    )
    payload = build_field_survey_render_payload(
        record=record,
        metrics=metrics,
        survey=survey,
        map_path=map_path,
    )
    from app.services.report_render.renderer import render_report_pdf

    pdf_path.write_bytes(render_report_pdf("field_survey", payload))
    for label, path in {
        "지도": map_path,
        "DOCX": docx_path,
        "PDF": pdf_path,
    }.items():
        if not path.is_file() or path.stat().st_size == 0:
            raise RuntimeError(f"{label} 생성 결과가 없습니다: {path}")
    return {
        "center_grid_id": record.center_grid_id,
        "year": record.year,
        "sido_name": record.sido_name,
        "sigungu_name": record.sigungu_name,
        "risk_score": round(float(metrics["risk_score"]), 2),
        "risk_grade": metrics["risk_grade"],
        "priority_score": round(float(metrics["priority_score"]), 2),
        "priority_grade": metrics["priority_grade"],
        "block_grid_ids": metrics["block_grid_ids"],
        "survey_datetime": survey.survey_datetime,
        "surveyors": survey.surveyors,
        "suspicious_count": survey.suspicious_count,
        "sample_count": survey.sample_count,
        "map_path": str(map_path),
        "docx_path": str(docx_path),
        "pdf_path": str(pdf_path),
    }
