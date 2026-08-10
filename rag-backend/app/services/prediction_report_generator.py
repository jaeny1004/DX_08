from __future__ import annotations

import io
import math
import os
import random
import re
import time
import zipfile
from dataclasses import dataclass
from datetime import datetime
from functools import lru_cache
from pathlib import Path
from typing import Any

import httpx
from docx import Document
from dotenv import load_dotenv
from PIL import Image, ImageChops, ImageDraw, ImageFont
from app.core.supabase_rest import create_client


BACKEND_ROOT = Path(__file__).resolve().parents[2]
# field_survey/control 생성기의 기존 import 계약을 유지하는 호환 alias.
# Vercel에서는 프로젝트와 백엔드 루트가 모두 /var/task이다.
PROJECT_ROOT = BACKEND_ROOT
TEMPLATE_PATH = (
    BACKEND_ROOT
    / "data"
    / "report_templates"
    / "[양식]소나무재선충병 발생 예측 보고서_빈양식.docx"
)
DEFAULT_OUTPUT_ROOT = (
    BACKEND_ROOT
    / "data"
    / "generated_drafts"
    / "prediction_template_test"
)

MAP_WIDTH = 1024
MAP_HEIGHT = 704
MAP_TILE_SIZE = 256
DEFAULT_ZOOM = 10
HISTORY_LAST_YEAR = 2021

DEFAULT_REPORT_TILE_URL = (
    "https://tile.openstreetmap.org/{z}/{x}/{y}.png"
)
DEFAULT_REPORT_TILE_ATTRIBUTION = "© OpenStreetMap contributors"

RISK_COLORS = {
    "매우 높음": (255, 59, 91, 200),
    "높음": (255, 145, 35, 195),
    "주의": (255, 207, 64, 190),
    "관찰": (57, 197, 105, 180),
    "낮음": (81, 160, 245, 115),
}


@dataclass
class ReportRecord:
    report_no: int
    year: int
    center_grid_id: int
    annual_count: int
    cumulative_count: int
    sido_name: str
    sigungu_name: str


@dataclass
class ReportSourceData:
    static: dict[str, Any]
    infection_positions: list[dict[str, Any]]
    stats: dict[str, Any]
    stats_year: int


def sanitize_filename(value: str) -> str:
    return re.sub(r'[\\/:*?"<>|]+', "_", value).strip()


def risk_grade(score: float) -> tuple[int, str]:
    if score >= 85:
        return 1, "매우 높음"
    if score >= 70:
        return 2, "높음"
    if score >= 55:
        return 3, "주의"
    if score >= 40:
        return 4, "관찰"
    return 5, "낮음"


def priority_grade(score: float) -> str:
    if score >= 85:
        return "최우선 예찰"
    if score >= 70:
        return "우선 예찰"
    if score >= 55:
        return "집중 관찰"
    if score >= 40:
        return "정기 관찰"
    return "일반 관리"


def _to_float(
    value: Any,
    default: float | None = None,
) -> float | None:
    try:
        if value is None:
            return default
        return float(value)
    except (TypeError, ValueError):
        return default


def _first_value(
    row: dict[str, Any],
    keys: Sequence[str],
    default: Any = None,
) -> Any:
    for key in keys:
        value = row.get(key)
        if value is None or str(value).strip() == "":
            continue
        return value
    return default


def _create_supabase_client() -> Any:
    load_dotenv(BACKEND_ROOT / ".env")
    url = os.getenv("SUPABASE_URL", "").strip()
    key = os.getenv("SUPABASE_KEY", "").strip()

    if not url or not key:
        raise RuntimeError(
            "SUPABASE_URL과 service_role SUPABASE_KEY가 필요합니다."
        )

    return create_client(url, key)


def _single_row(
    rows: list[dict[str, Any]],
    *,
    table: str,
    grid_id: int,
) -> dict[str, Any]:
    if not rows:
        raise ValueError(
            f"격자 {grid_id}는 신규 확산위험 후보가 아니거나 "
            f"{table} 사전계산 대상에 없습니다."
        )
    if len(rows) != 1:
        raise RuntimeError(
            f"{table}에서 격자 {grid_id}가 {len(rows)}건 조회됐습니다."
        )
    return dict(rows[0])


def load_report_source_data(
    center_grid_id: int,
    year: int,
    *,
    client: Any | None = None,
) -> ReportSourceData:
    client = client or _create_supabase_client()
    static_response = (
        client.table("prediction_grid_static")
        .select(
            "grid_id,risk_candidate_flag,center_point_4326,"
            "center_point_5186,"
            "cell_geometry_4326,block_geometry_4326,block_grid_ids,"
            "nearby_infection_grid_ids,sido_name,sigungu_name,"
            "block_cell_geometries_4326,block_pine_mean,"
            "block_elevation_mean,block_slope_mean"
        )
        .eq("grid_id", center_grid_id)
        .limit(2)
        .execute()
    )
    static = _single_row(
        list(static_response.data or []),
        table="prediction_grid_static",
        grid_id=center_grid_id,
    )

    if static.get("risk_candidate_flag") is not True:
        raise ValueError(
            f"격자 {center_grid_id}는 신규 확산위험 후보가 아닙니다."
        )

    required_static = [
        "center_point_4326",
        "center_point_5186",
        "cell_geometry_4326",
        "block_geometry_4326",
        "block_grid_ids",
        "block_cell_geometries_4326",
        "block_pine_mean",
        "block_elevation_mean",
        "block_slope_mean",
    ]
    missing = [key for key in required_static if static.get(key) is None]
    if missing:
        raise RuntimeError(
            f"격자 {center_grid_id}의 사전계산 필드가 비어 있습니다: {missing}"
        )

    block_ids = [int(value) for value in static["block_grid_ids"]]
    block_geometries = list(static["block_cell_geometries_4326"])
    if len(block_ids) != len(block_geometries):
        raise RuntimeError(
            f"격자 {center_grid_id}의 block_grid_ids와 geometry 배열 길이가 "
            f"다릅니다: {len(block_ids)} != {len(block_geometries)}"
        )
    if center_grid_id not in block_ids:
        raise RuntimeError(
            f"격자 {center_grid_id}의 block_grid_ids에 중심 격자가 없습니다."
        )

    nearby_ids = [
        int(value) for value in static.get("nearby_infection_grid_ids") or []
    ]
    infection_response = client.rpc(
        "get_infection_grid_positions",
        {"p_grid_ids": nearby_ids},
    ).execute()
    infection_payload = list(infection_response.data or [])
    if len(infection_payload) != 1:
        raise RuntimeError(
            "get_infection_grid_positions RPC가 단일 결과 행을 "
            f"반환하지 않았습니다: {len(infection_payload)}건"
        )
    infection_positions = [
        dict(row) for row in (infection_payload[0].get("positions") or [])
    ]
    infection_positions.sort(key=lambda row: int(row["grid_id"]))

    found_ids = {int(row["grid_id"]) for row in infection_positions}
    missing_infection_ids = sorted(set(nearby_ids) - found_ids)
    if missing_infection_ids:
        raise RuntimeError(
            f"격자 {center_grid_id}의 14km 감염 이력 위치 "
            f"{len(missing_infection_ids)}건을 찾지 못했습니다: "
            f"{missing_infection_ids[:20]}"
        )
    stats_year = year if 2016 <= year <= HISTORY_LAST_YEAR else HISTORY_LAST_YEAR
    stats_response = (
        client.table("grid_infection_stats")
        .select("*")
        .eq("grid_id", center_grid_id)
        .eq("year", stats_year)
        .execute()
    )
    stats = _single_row(
        list(stats_response.data or []),
        table=f"grid_infection_stats(year={stats_year})",
        grid_id=center_grid_id,
    )
    return ReportSourceData(
        static=static,
        infection_positions=infection_positions,
        stats=stats,
        stats_year=stats_year,
    )


def calculate_metrics(
    record: ReportRecord,
    source: ReportSourceData,
) -> dict[str, Any]:
    static = source.static
    stats = source.stats
    historical_year = 2016 <= record.year <= HISTORY_LAST_YEAR
    block_annual = int(stats["block_annual_count"]) if historical_year else 0
    active_count = (
        int(stats["block_active_grid_count"]) if historical_year else 0
    )
    block_cumulative = int(stats["block_cumulative_count"])
    pine_mean = float(static["block_pine_mean"])
    elevation_mean = float(static["block_elevation_mean"])
    slope_mean = float(static["block_slope_mean"])

    infection_pressure = min(
        98.0,
        25.0
        + 10.5 * math.log1p(max(block_cumulative, 0))
        + 4.5 * math.log1p(max(block_annual, 0)),
    )
    access_score = max(
        20.0,
        min(92.0, 88.0 - slope_mean * 1.15 - elevation_mean * 0.018),
    )
    road_distance = max(
        80.0,
        min(
            1800.0,
            90.0 + slope_mean * 24.0 + elevation_mean * 0.08,
        ),
    )
    risk_score = min(
        98.5,
        24.0
        + infection_pressure * 0.39
        + pine_mean * 0.28
        + (100.0 - access_score) * 0.10
        + math.log1p(max(record.annual_count, 1)) * 7.5,
    )
    stage, grade = risk_grade(risk_score)
    priority_score = min(
        98.0,
        risk_score * 0.72
        + infection_pressure * 0.20
        + access_score * 0.08,
    )

    if road_distance < 350:
        road_type = "주요 도로 인접 구간"
    elif road_distance < 800:
        road_type = "일반도로 인접 구간"
    elif road_distance < 1400:
        road_type = "산림 접근 구간"
    else:
        road_type = "산림 내부 접근 구간"

    return {
        "block_count": len(static["block_grid_ids"]),
        "block_grid_ids": [int(value) for value in static["block_grid_ids"]],
        "block_annual": block_annual,
        "block_cumulative": block_cumulative,
        "active_count": active_count,
        "pine_mean": pine_mean,
        "elevation_mean": elevation_mean,
        "slope_mean": slope_mean,
        "infection_pressure": infection_pressure,
        "access_score": access_score,
        "road_distance": road_distance,
        "road_type": road_type,
        "environment_warning": (
            "해당" if slope_mean >= 25 or elevation_mean >= 450 else "비해당"
        ),
        "risk_score": risk_score,
        "risk_stage": stage,
        "risk_grade": grade,
        "priority_score": priority_score,
        "priority_grade": priority_grade(priority_score),
    }


def apply_candidate_metrics(
    metrics: dict[str, Any],
    candidate: dict[str, Any] | None,
) -> dict[str, Any]:
    result = dict(metrics)
    candidate = candidate or {}

    risk_score = _to_float(
        _first_value(
            candidate,
            [
                "risk_score",
                "prediction_score",
                "pred_score",
                "ensemble_score",
                "risk_probability",
            ],
        )
    )
    if risk_score is not None:
        if 0.0 <= risk_score <= 1.0:
            risk_score *= 100.0
        risk_score = max(0.0, min(100.0, risk_score))
        stage, calculated_grade = risk_grade(risk_score)
        result.update(
            risk_score=risk_score,
            risk_stage=stage,
            risk_grade=str(
                _first_value(
                    candidate,
                    ["risk_grade", "risk_label", "risk_level"],
                    calculated_grade,
                )
            ),
        )

    priority_score = _to_float(
        _first_value(
            candidate,
            [
                "priority_score",
                "field_priority_score_v3",
                "field_priority_score",
                "survey_priority_score",
            ],
        )
    )
    if priority_score is not None:
        if 0.0 <= priority_score <= 1.0:
            priority_score *= 100.0
        priority_score = max(0.0, min(100.0, priority_score))
        result.update(
            priority_score=priority_score,
            priority_grade=str(
                _first_value(
                    candidate,
                    [
                        "priority_grade",
                        "priority_label",
                        "priority_stage_label",
                        "field_priority_label",
                    ],
                    priority_grade(priority_score),
                )
            ),
        )

    overrides = {
        "infection_pressure": [
            "infection_pressure",
            "recent_infection_pressure",
            "infection_pressure_score",
        ],
        "access_score": [
            "access_score",
            "access_score_v3",
            "accessibility_score",
        ],
        "road_distance": [
            "road_distance",
            "road_distance_m",
            "distance_to_road_m",
            "distance_to_nearest_road_m_v3",
            "nearest_road_distance_m",
        ],
    }
    for target, keys in overrides.items():
        value = _to_float(_first_value(candidate, keys))
        if value is None:
            continue
        if target != "road_distance" and 0.0 <= value <= 1.0:
            value *= 100.0
        result[target] = (
            max(0.0, value)
            if target == "road_distance"
            else max(0.0, min(100.0, value))
        )

    return result


def lonlat_to_world_pixel(
    lon: float,
    lat: float,
    zoom: int,
) -> tuple[float, float]:
    lat = max(min(lat, 85.05112878), -85.05112878)
    scale = 256.0 * (2**zoom)
    x = (lon + 180.0) / 360.0 * scale
    sin_lat = math.sin(math.radians(lat))
    y = (
        0.5
        - math.log((1 + sin_lat) / (1 - sin_lat)) / (4 * math.pi)
    ) * scale
    return x, y


def geometry_to_image_rings(
    geometry: dict[str, Any],
    center_lon: float,
    center_lat: float,
    zoom: int,
    width: int = MAP_WIDTH,
    height: int = MAP_HEIGHT,
) -> list[list[tuple[float, float]]]:
    center_x, center_y = lonlat_to_world_pixel(
        center_lon,
        center_lat,
        zoom,
    )
    geometry_type = geometry.get("type")
    coordinates = geometry.get("coordinates") or []
    if geometry_type == "Polygon":
        source_rings = coordinates[:1]
    elif geometry_type == "MultiPolygon":
        source_rings = [
            polygon[0] for polygon in coordinates if polygon
        ]
    else:
        raise ValueError(f"지원하지 않는 GeoJSON geometry: {geometry_type}")

    result: list[list[tuple[float, float]]] = []
    for ring in source_rings:
        points = []
        for lon, lat, *_ in ring:
            world_x, world_y = lonlat_to_world_pixel(
                float(lon),
                float(lat),
                zoom,
            )
            points.append(
                (
                    width / 2.0 + (world_x - center_x),
                    height / 2.0 + (world_y - center_y),
                )
            )
        result.append(points)
    return result


def local_risk_grade(annual: int, cumulative: int) -> str:
    score = (
        30.0
        + 12.0 * math.log1p(max(annual, 0))
        + 8.5 * math.log1p(max(cumulative, 0))
    )
    return risk_grade(min(score, 98.0))[1]


def request_vworld_map(
    api_key: str,
    domain: str,
    center_lon: float,
    center_lat: float,
    zoom: int,
    basemap: str,
    width: int = MAP_WIDTH,
    height: int = MAP_HEIGHT,
) -> Image.Image:
    endpoint = "https://api.vworld.kr/req/image"
    common = {
        "service": "image",
        "request": "GetMap",
        "version": "2.0",
        "key": api_key,
        "domain": domain,
        "format": "png",
        "transparent": "false",
        "center": f"{center_lon:.8f},{center_lat:.8f}",
        "crs": "EPSG:4326",
        "zoom": str(zoom),
        "size": f"{width},{height}",
        "basemap": basemap,
    }
    attempts = [common, {**common, "request": "getmap"}]
    last_error: Exception | None = None

    for params in attempts:
        try:
            response = httpx.get(
                endpoint,
                params=params,
                timeout=httpx.Timeout(10.0, connect=5.0),
                headers={
                    "Accept": "image/png,image/*;q=0.9,*/*;q=0.1",
                    "User-Agent": "PineWiltReportGenerator/1.0",
                    "Connection": "close",
                },
                follow_redirects=True,
            )
            response.raise_for_status()
            if "image" not in response.headers.get(
                "content-type",
                "",
            ).lower():
                raise RuntimeError(
                    "VWorld가 이미지 대신 다음 내용을 반환했습니다: "
                    f"{response.text[:500]}"
                )
            image = Image.open(io.BytesIO(response.content)).convert("RGBA")
            if image.size != (width, height):
                image = image.resize(
                    (width, height),
                    Image.Resampling.LANCZOS,
                )
            return image
        except Exception as exc:
            last_error = exc
            time.sleep(1)

    raise RuntimeError(
        "VWorld 지도 이미지를 받지 못했습니다. "
        f"마지막 오류: {last_error}"
    )


def request_xyz_tile_map(
    center_lon: float,
    center_lat: float,
    zoom: int,
    *,
    tile_url: str = DEFAULT_REPORT_TILE_URL,
    width: int = MAP_WIDTH,
    height: int = MAP_HEIGHT,
    referer: str | None = None,
) -> Image.Image:
    """Leaflet과 같은 XYZ 타일을 합쳐 보고서용 실제 배경지도를 만든다.

    한 장이라도 내려받지 못하면 예외를 발생시킨다. 호출부에서 다음 지도
    공급자 또는 완전 로컬 배경으로 전환하므로, 일부만 빈 지도가 PDF에
    들어가는 상황을 방지한다.
    """
    required_tokens = ("{z}", "{x}", "{y}")
    if not all(token in tile_url for token in required_tokens):
        raise ValueError(
            "REPORT_MAP_TILE_URL에는 {z}, {x}, {y}가 모두 필요합니다."
        )

    center_x, center_y = lonlat_to_world_pixel(
        center_lon,
        center_lat,
        zoom,
    )
    left = center_x - width / 2.0
    top = center_y - height / 2.0
    right = left + width
    bottom = top + height

    min_tile_x = math.floor(left / MAP_TILE_SIZE)
    max_tile_x = math.floor((right - 1) / MAP_TILE_SIZE)
    min_tile_y = math.floor(top / MAP_TILE_SIZE)
    max_tile_y = math.floor((bottom - 1) / MAP_TILE_SIZE)
    tile_count = 2**zoom

    canvas = Image.new("RGBA", (width, height), (241, 245, 249, 255))
    headers = {
        "Accept": "image/png,image/jpeg,image/*;q=0.9,*/*;q=0.1",
        "User-Agent": "PineWiltAdministrativeReport/1.0",
    }
    if referer:
        headers["Referer"] = referer

    with httpx.Client(
        timeout=httpx.Timeout(4.0, connect=3.0),
        headers=headers,
        follow_redirects=True,
    ) as client:
        for tile_y in range(min_tile_y, max_tile_y + 1):
            if tile_y < 0 or tile_y >= tile_count:
                raise RuntimeError(
                    f"지도 타일 Y 좌표가 범위를 벗어났습니다: {tile_y}"
                )
            for tile_x in range(min_tile_x, max_tile_x + 1):
                wrapped_x = tile_x % tile_count
                url = (
                    tile_url.replace("{s}", "a")
                    .replace("{z}", str(zoom))
                    .replace("{x}", str(wrapped_x))
                    .replace("{y}", str(tile_y))
                )
                response = client.get(url)
                response.raise_for_status()
                if "image" not in response.headers.get(
                    "content-type",
                    "",
                ).lower():
                    raise RuntimeError(
                        "지도 타일 서버가 이미지가 아닌 응답을 반환했습니다: "
                        f"{response.text[:200]}"
                    )
                tile = Image.open(io.BytesIO(response.content)).convert("RGBA")
                if tile.size != (MAP_TILE_SIZE, MAP_TILE_SIZE):
                    tile = tile.resize(
                        (MAP_TILE_SIZE, MAP_TILE_SIZE),
                        Image.Resampling.LANCZOS,
                    )
                paste_x = round(tile_x * MAP_TILE_SIZE - left)
                paste_y = round(tile_y * MAP_TILE_SIZE - top)
                canvas.paste(tile, (paste_x, paste_y))

    return canvas


def create_fallback_map_background(
    center_lon: float,
    center_lat: float,
    width: int = MAP_WIDTH,
    height: int = MAP_HEIGHT,
) -> Image.Image:
    """VWorld 장애 시에도 보고서를 만들 수 있는 로컬 분석 배경을 생성한다."""
    image = Image.new("RGBA", (width, height), (241, 245, 249, 255))
    draw = ImageDraw.Draw(image, "RGBA")

    # 지도처럼 위치를 가늠할 수 있는 좌표 격자
    grid_step = 64
    for x in range(0, width + 1, grid_step):
        draw.line((x, 0, x, height), fill=(148, 163, 184, 58), width=1)
    for y in range(0, height + 1, grid_step):
        draw.line((0, y, width, y), fill=(148, 163, 184, 58), width=1)

    # 외부 타일 없이도 배경과 분석 격자가 구분되도록 단순 지형 요소를 표시
    river_points = [
        (-40, int(height * 0.18)),
        (int(width * 0.18), int(height * 0.28)),
        (int(width * 0.38), int(height * 0.22)),
        (int(width * 0.58), int(height * 0.42)),
        (int(width * 0.78), int(height * 0.38)),
        (width + 40, int(height * 0.52)),
    ]
    draw.line(river_points, fill=(125, 211, 252, 95), width=42, joint="curve")
    draw.line(river_points, fill=(56, 189, 248, 115), width=5, joint="curve")

    road_points = [
        (int(width * 0.05), height + 30),
        (int(width * 0.24), int(height * 0.72)),
        (int(width * 0.42), int(height * 0.62)),
        (int(width * 0.63), int(height * 0.44)),
        (int(width * 0.82), int(height * 0.16)),
        (width + 20, int(height * 0.08)),
    ]
    draw.line(road_points, fill=(255, 255, 255, 235), width=18, joint="curve")
    draw.line(road_points, fill=(245, 158, 11, 155), width=4, joint="curve")

    center_x, center_y = width // 2, height // 2
    draw.line((center_x - 18, center_y, center_x + 18, center_y), fill=(15, 23, 42, 150), width=2)
    draw.line((center_x, center_y - 18, center_x, center_y + 18), fill=(15, 23, 42, 150), width=2)

    small_font = find_font(16, bold=False)
    label = (
        "외부 배경지도 연결 불가 · 로컬 격자 분석 배경 "
        f"({center_lat:.5f}, {center_lon:.5f})"
    )
    draw.rounded_rectangle(
        (18, 18, 610, 50),
        radius=8,
        fill=(255, 255, 255, 225),
        outline=(148, 163, 184, 180),
        width=1,
    )
    draw.text((28, 25), label, font=small_font, fill=(71, 85, 105, 255))
    return image


@lru_cache(maxsize=32)
def find_font(
    size: int,
    bold: bool = False,
) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    bundled_font_dir = BACKEND_ROOT / "assets" / "fonts"
    candidates = [
        str(
            bundled_font_dir
            / (
                "NotoSansKR-Bold.ttf"
                if bold
                else "NotoSansKR-Regular.ttf"
            )
        ),
        str(
            bundled_font_dir
            / (
                "NotoSansKR-Bold.otf"
                if bold
                else "NotoSansKR-Regular.otf"
            )
        ),
        str(
            bundled_font_dir
            / (
                "NotoSansKR-Bold.woff"
                if bold
                else "NotoSansKR-Regular.woff"
            )
        ),
        (
            "/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc"
            if bold
            else "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"
        ),
        (
            "/usr/share/fonts/truetype/noto/NotoSansKR-Bold.ttf"
            if bold
            else "/usr/share/fonts/truetype/noto/NotoSansKR-Regular.ttf"
        ),
        (
            "/usr/share/fonts/truetype/nanum/NanumGothicBold.ttf"
            if bold
            else "/usr/share/fonts/truetype/nanum/NanumGothic.ttf"
        ),
        (
            "/usr/share/fonts/truetype/nanum/NanumSquareRoundB.ttf"
            if bold
            else "/usr/share/fonts/truetype/nanum/NanumSquareRoundR.ttf"
        ),
        (
            "/usr/share/fonts/truetype/unfonts-core/UnDotumBold.ttf"
            if bold
            else "/usr/share/fonts/truetype/unfonts-core/UnDotum.ttf"
        ),
    ]
    for path in candidates:
        if Path(path).exists():
            return ImageFont.truetype(path, size=size)
    return ImageFont.load_default()


def render_vworld_overlay(
    background: Image.Image,
    record: ReportRecord,
    source: ReportSourceData,
    zoom: int,
    background_source: str = "VWorld",
) -> tuple[Image.Image, dict[str, Any]]:
    static = source.static
    center_coordinates = static["center_point_4326"]["coordinates"]
    center_lon = float(center_coordinates[0])
    center_lat = float(center_coordinates[1])
    image = background.convert("RGBA")
    if image.size != (MAP_WIDTH, MAP_HEIGHT):
        image = image.resize(
            (MAP_WIDTH, MAP_HEIGHT),
            Image.Resampling.LANCZOS,
        )
    overlay = Image.new("RGBA", image.size, (0, 0, 0, 0))
    draw = ImageDraw.Draw(overlay, "RGBA")
    annual_column = f"infection_count_{record.year}"
    trace: dict[str, Any] = {
        "center": [center_lon, center_lat],
        "infection_polygons": {},
        "block_polygons": {},
        "block_outline": [],
    }

    for row in source.infection_positions:
        annual = int(row.get(annual_column, 0) or 0)
        cumulative = int(row.get("infection_count_2016_2021", 0) or 0)
        if annual <= 0 and cumulative <= 0:
            continue
        grade = local_risk_grade(annual, cumulative)
        rings = geometry_to_image_rings(
            row["geometry_4326"],
            center_lon,
            center_lat,
            zoom,
        )
        trace["infection_polygons"][str(row["grid_id"])] = rings
        for points in rings:
            if points and any(
                -100 <= x <= MAP_WIDTH + 100
                and -100 <= y <= MAP_HEIGHT + 100
                for x, y in points
            ):
                draw.polygon(
                    points,
                    fill=RISK_COLORS[grade],
                    outline=(255, 255, 255, 120),
                    width=1,
                )

    block_ids = [int(value) for value in static["block_grid_ids"]]
    for grid_id, geometry in zip(
        block_ids,
        static["block_cell_geometries_4326"],
    ):
        rings = geometry_to_image_rings(
            geometry,
            center_lon,
            center_lat,
            zoom,
        )
        trace["block_polygons"][str(grid_id)] = rings
        for points in rings:
            if grid_id == record.center_grid_id:
                draw.polygon(
                    points,
                    fill=(255, 56, 56, 75),
                    outline=(160, 0, 0, 255),
                    width=5,
                )
            else:
                draw.line(
                    points,
                    fill=(15, 75, 180, 240),
                    width=4,
                    joint="curve",
                )

    block_outline_rings = geometry_to_image_rings(
        static["block_geometry_4326"],
        center_lon,
        center_lat,
        zoom,
    )
    trace["block_outline"] = block_outline_rings
    for points in block_outline_rings:
        draw.line(
            points,
            fill=(0, 45, 130, 255),
            width=6,
            joint="curve",
        )

    title_font = find_font(32, bold=True)
    label_font = find_font(24, bold=True)
    small_font = find_font(18, bold=False)
    title = (
        f"{record.sido_name} {record.sigungu_name} / "
        f"{record.year}년 위험격자 분포"
    )
    title_box = draw.textbbox((0, 0), title, font=title_font)
    title_width = title_box[2] - title_box[0]
    draw.rounded_rectangle(
        (
            MAP_WIDTH / 2 - title_width / 2 - 18,
            16,
            MAP_WIDTH / 2 + title_width / 2 + 18,
            64,
        ),
        radius=10,
        fill=(255, 255, 255, 225),
        outline=(160, 160, 160, 220),
        width=2,
    )
    draw.text(
        (MAP_WIDTH / 2 - title_width / 2, 23),
        title,
        font=title_font,
        fill=(20, 20, 20, 255),
    )

    center_label = f"중심 격자 {record.center_grid_id}"
    label_box = draw.textbbox((0, 0), center_label, font=label_font)
    label_width = label_box[2] - label_box[0]
    label_x = MAP_WIDTH / 2 + 20
    label_y = MAP_HEIGHT / 2 - 65
    draw.rounded_rectangle(
        (
            label_x - 8,
            label_y - 6,
            label_x + label_width + 12,
            label_y + 34,
        ),
        radius=8,
        fill=(255, 255, 255, 235),
        outline=(170, 170, 170, 255),
        width=2,
    )
    draw.text(
        (label_x, label_y),
        center_label,
        font=label_font,
        fill=(20, 20, 20, 255),
    )

    legend_x = 24
    legend_y = MAP_HEIGHT - 165
    draw.rounded_rectangle(
        (legend_x, legend_y, legend_x + 420, legend_y + 135),
        radius=12,
        fill=(255, 255, 255, 235),
        outline=(175, 175, 175, 230),
        width=2,
    )
    legend_items = [
        ("매우 높음", RISK_COLORS["매우 높음"]),
        ("높음", RISK_COLORS["높음"]),
        ("주의", RISK_COLORS["주의"]),
        ("관찰", RISK_COLORS["관찰"]),
    ]
    for index, (name, color) in enumerate(legend_items):
        column = index % 2
        row_number = index // 2
        x = legend_x + 20 + column * 195
        y = legend_y + 17 + row_number * 45
        draw.rectangle(
            (x, y, x + 34, y + 25),
            fill=color,
            outline=(255, 255, 255, 220),
            width=1,
        )
        draw.text(
            (x + 45, y - 1),
            name,
            font=small_font,
            fill=(20, 20, 20, 255),
        )
    draw.rectangle(
        (legend_x + 20, legend_y + 99, legend_x + 54, legend_y + 124),
        fill=(255, 56, 56, 70),
        outline=(160, 0, 0, 255),
        width=4,
    )
    draw.text(
        (legend_x + 65, legend_y + 97),
        "중심 격자",
        font=small_font,
        fill=(20, 20, 20, 255),
    )
    draw.rectangle(
        (legend_x + 215, legend_y + 99, legend_x + 249, legend_y + 124),
        fill=(255, 255, 255, 30),
        outline=(15, 75, 180, 255),
        width=4,
    )
    draw.text(
        (legend_x + 260, legend_y + 97),
        "주변 8개 격자",
        font=small_font,
        fill=(20, 20, 20, 255),
    )

    source_text = (
        f"배경지도: {background_source} | "
        "격자: 2016~2021 감염 발생 이력"
    )
    source_box = draw.textbbox((0, 0), source_text, font=small_font)
    source_width = source_box[2] - source_box[0]
    draw.rounded_rectangle(
        (
            MAP_WIDTH - source_width - 35,
            MAP_HEIGHT - 38,
            MAP_WIDTH - 15,
            MAP_HEIGHT - 8,
        ),
        radius=6,
        fill=(255, 255, 255, 220),
    )
    draw.text(
        (MAP_WIDTH - source_width - 25, MAP_HEIGHT - 35),
        source_text,
        font=small_font,
        fill=(55, 55, 55, 255),
    )
    return Image.alpha_composite(image, overlay).convert("RGB"), trace


def _referer_url(domain: str | None) -> str | None:
    """VWORLD_API_DOMAIN을 Referer 헤더에 쓸 수 있는 형태로 정규화한다.

    환경변수에는 스킴 없이 도메인만 들어 있는데(예: example.vercel.app),
    그대로 Referer로 보내면 VWorld WMTS가 500을 돌려준다.
    그러면 배경지도가 조용히 OSM으로 넘어가 보고서에 다른 지도가 실린다.
    """
    value = (domain or "").strip()
    if not value:
        return None
    if value.startswith(("http://", "https://")):
        return value
    return f"https://{value}"


def build_vworld_overlay_map(
    output_path: Path,
    record: ReportRecord,
    source: ReportSourceData,
    api_key: str,
    domain: str,
    zoom: int,
    basemap: str,
    *,
    background: Image.Image | None = None,
) -> dict[str, Any]:
    center = source.static["center_point_4326"]["coordinates"]
    center_lon = float(center[0])
    center_lat = float(center[1])
    background_source = "제공된 배경지도"

    if background is None:
        errors: list[str] = []
        remote_default = "false" if os.getenv("VERCEL") else "true"
        vworld_enabled = os.getenv(
            "VWORLD_REMOTE_ENABLED",
            remote_default,
        ).strip().lower() in {"1", "true", "yes", "on"}

        # 1순위: 대시보드가 사용하는 것과 같은 VWorld Base WMTS 타일.
        if vworld_enabled and api_key:
            try:
                background = request_xyz_tile_map(
                    center_lon,
                    center_lat,
                    zoom,
                    tile_url=(
                        "https://api.vworld.kr/req/wmts/1.0.0/"
                        f"{api_key}/Base/{{z}}/{{y}}/{{x}}.png"
                    ),
                    referer=_referer_url(domain),
                )
                background_source = "© VWorld"
            except Exception as exc:
                errors.append(f"VWorld WMTS: {type(exc).__name__}: {exc}")

        # 2순위: 프론트엔드의 VWorld 키 미설정 시와 동일한 OSM 실제 지도.
        if background is None:
            tile_url = os.getenv(
                "REPORT_MAP_TILE_URL",
                DEFAULT_REPORT_TILE_URL,
            ).strip() or DEFAULT_REPORT_TILE_URL
            tile_attribution = os.getenv(
                "REPORT_MAP_TILE_ATTRIBUTION",
                DEFAULT_REPORT_TILE_ATTRIBUTION,
            ).strip() or DEFAULT_REPORT_TILE_ATTRIBUTION
            try:
                background = request_xyz_tile_map(
                    center_lon,
                    center_lat,
                    zoom,
                    tile_url=tile_url,
                )
                background_source = tile_attribution
            except Exception as exc:
                errors.append(f"XYZ 실제 지도: {type(exc).__name__}: {exc}")

        # 3순위: 타일 공급자가 모두 실패했을 때 기존 VWorld 정적 API도 시도한다.
        if background is None and vworld_enabled and api_key and domain:
            try:
                background = request_vworld_map(
                    api_key=api_key,
                    domain=domain,
                    center_lon=center_lon,
                    center_lat=center_lat,
                    zoom=zoom,
                    basemap=basemap,
                )
                background_source = "© VWorld"
            except Exception as exc:
                errors.append(f"VWorld 정적 지도: {type(exc).__name__}: {exc}")

        # 최종 안전장치: 모든 외부 지도 서비스가 실패해도 문서 생성은 계속한다.
        if background is None:
            print(
                "[prediction-report] 모든 실제 배경지도 요청 실패. "
                "로컬 대체 배경을 사용합니다: " + " | ".join(errors)
            )
            background = create_fallback_map_background(
                center_lon=center_lon,
                center_lat=center_lat,
            )
            background_source = "로컬 대체 배경"

    result, trace = render_vworld_overlay(
        background,
        record,
        source,
        zoom,
        background_source=background_source,
    )
    trace["background_source"] = background_source
    output_path.parent.mkdir(parents=True, exist_ok=True)
    result.save(output_path, quality=95)
    return trace


def replace_in_paragraph(
    paragraph: Any,
    replacements: dict[str, str],
) -> None:
    full = "".join(run.text for run in paragraph.runs)
    if not full:
        return
    new = full
    for old, value in replacements.items():
        new = new.replace(old, str(value))
    if new == full:
        return
    if paragraph.runs:
        paragraph.runs[0].text = new
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.text = new


def replace_everywhere(
    document: Document,
    replacements: dict[str, str],
) -> None:
    for paragraph in document.paragraphs:
        replace_in_paragraph(paragraph, replacements)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    replace_in_paragraph(paragraph, replacements)
    for section in document.sections:
        for paragraph in section.header.paragraphs:
            replace_in_paragraph(paragraph, replacements)
        for paragraph in section.footer.paragraphs:
            replace_in_paragraph(paragraph, replacements)


def set_paragraph(
    document: Document,
    prefix: str,
    value: str,
) -> None:
    for paragraph in document.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            if paragraph.runs:
                paragraph.runs[0].text = value
                for run in paragraph.runs[1:]:
                    run.text = ""
            else:
                paragraph.text = value
            return
    raise RuntimeError(f"템플릿에서 문단을 찾지 못했습니다: {prefix}")


def replace_docx_media_images(
    docx_path: Path,
    replacements: dict[str, Path],
) -> None:
    temporary_path = docx_path.with_suffix(".tmp.docx")
    replaced: set[str] = set()
    with zipfile.ZipFile(docx_path, "r") as source, zipfile.ZipFile(
        temporary_path,
        "w",
        zipfile.ZIP_DEFLATED,
    ) as target:
        for item in source.infolist():
            replacement = replacements.get(item.filename)
            if replacement is not None:
                target.writestr(item, replacement.read_bytes())
                replaced.add(item.filename)
            else:
                target.writestr(item, source.read(item.filename))
    missing = set(replacements) - replaced
    if missing:
        temporary_path.unlink(missing_ok=True)
        raise RuntimeError(
            f"DOCX 안에서 교체할 media를 찾지 못했습니다: {sorted(missing)}"
        )
    temporary_path.replace(docx_path)


def extract_template_media(
    template_path: Path,
    media_name: str,
) -> Image.Image:
    with zipfile.ZipFile(template_path, "r") as archive:
        try:
            data = archive.read(media_name)
        except KeyError as exc:
            raise RuntimeError(
                f"템플릿 안에서 {media_name}을 찾지 못했습니다."
            ) from exc
    return Image.open(io.BytesIO(data)).convert("RGB")


def draw_text_in_box(
    draw: ImageDraw.ImageDraw,
    box_xy: tuple[int, int, int, int],
    value: str,
    font: ImageFont.ImageFont,
    fill: tuple[int, int, int] = (25, 25, 25),
    max_lines: int = 3,
) -> None:
    left, top, right, bottom = box_xy
    max_width = max(8, right - left - 6)
    words = str(value).replace("\n", " \n ").split()
    lines: list[str] = []
    current = ""
    for word in words:
        if word == "\n":
            if current:
                lines.append(current)
                current = ""
            continue
        candidate = word if not current else f"{current} {word}"
        box = draw.textbbox((0, 0), candidate, font=font)
        if box[2] - box[0] <= max_width:
            current = candidate
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    if not lines:
        lines = [""]
    if len(lines) > max_lines:
        lines = lines[:max_lines]
        last = lines[-1]
        while last:
            candidate = last + "…"
            box = draw.textbbox((0, 0), candidate, font=font)
            if box[2] - box[0] <= max_width:
                lines[-1] = candidate
                break
            last = last[:-1]
    heights = []
    for line in lines:
        box = draw.textbbox((0, 0), line, font=font)
        heights.append(max(1, box[3] - box[1]))
    spacing = 2
    total_height = sum(heights) + spacing * (len(lines) - 1)
    y = top + max(0, (bottom - top - total_height) / 2)
    for line, height in zip(lines, heights):
        box = draw.textbbox((0, 0), line, font=font)
        width = box[2] - box[0]
        x = left + max(0, (right - left - width) / 2)
        draw.text((x, y), line, font=font, fill=fill)
        y += height + spacing


def validate_prediction_template(template_path: Path) -> None:
    expected = {
        "word/media/image1.png": (1200, 1240),
        "word/media/image2.png": (826, 1264),
        "word/media/image3.png": (826, 1264),
    }
    with zipfile.ZipFile(template_path, "r") as archive:
        names = set(archive.namelist())
        missing = set(expected) - names
        if missing:
            raise RuntimeError(
                "발생 예측 보고서 양식의 이미지 구성이 예상과 다릅니다. "
                f"누락 media: {sorted(missing)}"
            )
        for name, expected_size in expected.items():
            image = Image.open(io.BytesIO(archive.read(name)))
            if image.size != expected_size:
                raise RuntimeError(
                    f"양식 media 크기가 다릅니다: {name} "
                    f"예상 {expected_size}, 실제 {image.size}"
                )


def verify_appendix_was_filled(
    original: Image.Image,
    filled: Image.Image,
    name: str,
) -> None:
    if original.size != filled.size:
        raise RuntimeError(f"{name} 이미지 크기가 원본과 달라졌습니다.")
    difference = ImageChops.difference(
        original.convert("RGB"),
        filled.convert("RGB"),
    )
    changed = sum(
        1
        for red, green, blue in difference.getdata()
        if red > 8 or green > 8 or blue > 8
    )
    if changed < 800:
        raise RuntimeError(f"{name} 별지에 입력 내용이 그려지지 않았습니다.")


def build_appendix_entries(
    record: ReportRecord,
    metrics: dict[str, Any],
) -> tuple[list[list[str]], list[list[str]]]:
    generator = random.Random(
        record.year * 1_000_003 + record.center_grid_id
    )
    names = [
        "김도윤",
        "박서준",
        "이현우",
        "최민재",
        "정하늘",
        "윤지호",
        "한예린",
        "송민수",
    ]
    reporter_types = ["주민신고", "순찰확인", "드론예찰", "산림관계자"]
    contents = [
        "소나무 수관 변색 및 잎 마름 관찰",
        "도로변 고사목과 부분 갈변목 확인",
        "산림 가장자리 고사 의심목 발견",
        "드론 영상에서 수관 이상 후보 확인",
    ]
    results = [
        "현장 확인 예정",
        "목편 시료 채취 후 검경 의뢰",
        "주변 격자 추가 예찰",
        "담당부서 전달 및 재방문 예정",
    ]
    report_rows: list[list[str]] = []
    row_count = 4 if metrics["risk_score"] >= 70 else 3
    adjacent = [
        grid_id
        for grid_id in metrics["block_grid_ids"]
        if grid_id != record.center_grid_id
    ]
    all_grids = [record.center_grid_id] + adjacent
    for index in range(row_count):
        month = 4 + ((record.report_no + index) % 5)
        day = 7 + ((record.center_grid_id + index * 5) % 20)
        grid_id = all_grids[index % len(all_grids)]
        report_rows.append(
            [
                reporter_types[index % len(reporter_types)],
                f"{record.year}.{month:02d}.{day:02d}",
                names[(record.report_no + index) % len(names)],
                f"{record.sigungu_name}\n격자 {grid_id}",
                contents[(record.center_grid_id + index) % len(contents)],
                results[(record.report_no + index) % len(results)],
                (
                    "010-****-"
                    f"{1000 + ((record.center_grid_id + index * 137) % 9000):04d}"
                ),
            ]
        )

    planned_area = min(
        225,
        max(
            100,
            int(
                metrics["block_count"] * 25
                - generator.randint(0, 35)
            ),
        ),
    )
    plan_rows = [
        [
            f"{record.sido_name}\n{record.sigungu_name}",
            f"중심격자\n{record.center_grid_id}\n3x3 권역",
            str(planned_area),
            (
                f"{record.year}.{5 + record.report_no % 4:02d}."
                f"{10 + record.report_no % 15:02d}"
            ),
            f"{record.sigungu_name}\n임시 이착륙장",
            f"{record.sigungu_name}\n산림부서",
            "산림주사",
            names[(record.report_no + 2) % len(names)],
            "산림청\n지원헬기",
        ],
        [
            f"{record.sido_name}\n{record.sigungu_name}",
            "인접 고위험\n격자권역",
            str(max(75, planned_area - generator.randint(20, 55))),
            (
                f"{record.year}.{6 + record.report_no % 3:02d}."
                f"{12 + record.report_no % 13:02d}"
            ),
            f"{record.sigungu_name}\n공공시설 인근",
            f"{record.sigungu_name}\n산림부서",
            "주무관",
            names[(record.report_no + 5) % len(names)],
            "임차\n중형헬기",
        ],
    ]
    return report_rows, plan_rows


def create_filled_appendix_images(
    template_path: Path,
    output_directory: Path,
    record: ReportRecord,
    metrics: dict[str, Any],
) -> tuple[Path, Path]:
    output_directory.mkdir(parents=True, exist_ok=True)
    report_rows, plan_rows = build_appendix_entries(record, metrics)

    report_original = extract_template_media(
        template_path,
        "word/media/image2.png",
    )
    report_image = report_original.copy()
    report_draw = ImageDraw.Draw(report_image)
    report_font = find_font(13, bold=False)
    report_x = [29, 141, 253, 364, 476, 592, 704, 815]
    report_top = 220
    report_height = 45
    for row_index, values in enumerate(report_rows):
        top = report_top + row_index * report_height
        bottom = top + report_height
        for column_index, value in enumerate(values):
            cell = (
                report_x[column_index] + 1,
                top + 1,
                report_x[column_index + 1] - 1,
                bottom - 1,
            )
            report_draw.rectangle(cell, fill=(255, 255, 255))
            draw_text_in_box(
                report_draw,
                (
                    report_x[column_index] + 2,
                    top + 2,
                    report_x[column_index + 1] - 2,
                    bottom - 2,
                ),
                value,
                report_font,
                max_lines=3,
            )
    verify_appendix_was_filled(
        report_original,
        report_image,
        "신고 접수·처리 대장",
    )
    report_path = output_directory / (
        f"{record.report_no:02d}_{record.year}_"
        f"{record.center_grid_id}_신고접수대장.png"
    )
    report_image.save(report_path, quality=95)

    plan_original = extract_template_media(
        template_path,
        "word/media/image3.png",
    )
    plan_image = plan_original.copy()
    plan_draw = ImageDraw.Draw(plan_image)
    plan_font = find_font(12, bold=False)
    plan_x = [24, 130, 210, 295, 375, 471, 554, 637, 721, 814]
    plan_top = 240
    plan_height = 48
    for row_index, values in enumerate(plan_rows):
        top = plan_top + row_index * plan_height
        bottom = top + plan_height
        for column_index, value in enumerate(values):
            cell = (
                plan_x[column_index] + 1,
                top + 1,
                plan_x[column_index + 1] - 1,
                bottom - 1,
            )
            plan_draw.rectangle(cell, fill=(255, 255, 255))
            draw_text_in_box(
                plan_draw,
                (
                    plan_x[column_index] + 2,
                    top + 2,
                    plan_x[column_index + 1] - 2,
                    bottom - 2,
                ),
                value,
                plan_font,
                max_lines=3,
            )
    verify_appendix_was_filled(
        plan_original,
        plan_image,
        "유인항공예찰 계획",
    )
    plan_path = output_directory / (
        f"{record.report_no:02d}_{record.year}_"
        f"{record.center_grid_id}_유인항공예찰계획.png"
    )
    plan_image.save(plan_path, quality=95)
    return report_path, plan_path


def create_docx(
    template_path: Path,
    output_path: Path,
    map_path: Path,
    appendix_directory: Path,
    record: ReportRecord,
    metrics: dict[str, Any],
    start_date: str | None = None,
    end_date: str | None = None,
) -> None:
    document = Document(template_path)
    region = f"{record.sido_name} {record.sigungu_name}"
    # 사용자가 입력한 기간을 쓴다. 예전에는 연도만 표시하고 시작일·종료일을
    # 버렸으며, 작성일마저 12월 임의 날짜로 지어내고 있었다.
    period = (
        f"{start_date} ~ {end_date}"
        if start_date and end_date
        else f"{record.year}년"
    )
    grid_ids = metrics["block_grid_ids"]
    adjacent_ids = [
        grid_id
        for grid_id in grid_ids
        if grid_id != record.center_grid_id
    ]
    while len(adjacent_ids) < 4:
        adjacent_ids.append(record.center_grid_id)
    replacements = {
        "[작성일]": datetime.now().strftime("%Y. %m. %d."),
        "-지역, 기간-": f"-{region}, {period}-",
        "[지역]": region,
        "[기간]": period,
        "[시작일]": start_date or f"{record.year}-01-01",
        "[종료일]": end_date or f"{record.year}-12-31",
        "[격자 ID]": str(record.center_grid_id),
        "[단계 수]": "5",
        "[단계]": str(metrics["risk_stage"]),
        "[등급]": metrics["risk_grade"],
        "[비율]": f'{metrics["pine_mean"]:.1f}',
        "[거리]": f'{metrics["road_distance"]:.0f}',
        "[도로 유형]": metrics["road_type"],
        "[해당 여부]": metrics["environment_warning"],
        "[방향·대상 지역]": "중심 격자 주변 3x3 예찰 검토권역",
        "[결과]": "높음" if metrics["risk_score"] >= 70 else "중간",
        "[방향]": "중심 격자 주변",
        "[인접 격자 1]": str(adjacent_ids[0]),
        "[인접 격자 2]": str(adjacent_ids[1]),
        "[인접 격자 3]": str(adjacent_ids[2]),
        "[인접 격자 4]": str(adjacent_ids[3]),
        "[대응 단계명]": "우선 예찰 검토",
        # 실행 일자도 임의 12월 날짜가 아니라 사용자가 지정한 종료일을 쓴다.
        "[일자]": end_date or f"{record.year}-12-31",
        "[대상 격자]": ", ".join(map(str, grid_ids)),
    }
    replace_everywhere(document, replacements)
    paragraphs = {
        "❍ (분석 배경)": (
            f"❍ (분석 배경) {record.year}년 감염 발생 이력과 500m 격자 기반 "
            "산림·지형정보를 종합하여 신규 확산위험 후보를 분석함"
        ),
        "❍ (분석 목적)": (
            f"❍ (분석 목적) {region} 내 중심 격자 {record.center_grid_id}와 "
            "주변 격자를 하나의 예찰 검토권역으로 설정하여 현장 확인 필요성과 "
            "예찰 우선순위를 판단"
        ),
        "❍ (활용 목적)": (
            "❍ (활용 목적) AI 위험도 분석 결과를 예찰 계획 수립, 현장 확인, "
            "후속 방제 검토 및 행정 보고에 활용"
        ),
        "❍ (분석 대상)": (
            f"❍ (분석 대상) {region} / {period} / 중심 격자 포함 "
            f"3x3 권역 {metrics['block_count']}개 격자"
        ),
        "❍ (위험 점수)": (
            f"❍ (위험 점수) 종합 위험도 스코어 {metrics['risk_score']:.1f}점 "
            f"(전체 5단계 중 {metrics['risk_stage']}단계 "
            f"‘{metrics['risk_grade']}’ 수준)"
        ),
        "― 최근 감염압력": (
            f"― 최근 감염압력 {metrics['infection_pressure']:.1f}점"
            f"(해당 연도 권역 발생 {metrics['block_annual']}건, "
            f"2016~2021 누적 {metrics['block_cumulative']}건 반영)"
        ),
        "― 소나무류 비율": (
            f"― 소나무류 비율 {metrics['pine_mean']:.1f}%"
        ),
        "❍ 예찰 우선순위": (
            f"❍ 예찰 우선순위 {metrics['priority_score']:.1f}점"
            f"({metrics['priority_grade']})"
        ),
        "❍ 접근성": f"❍ 접근성 {metrics['access_score']:.1f}점",
        "― 도로까지의 거리": (
            f"― 도로까지의 거리 {metrics['road_distance']:.0f}m"
            f"({metrics['road_type']})"
        ),
        "― 환경주의": f"― 환경주의 {metrics['environment_warning']}",
        "❍ 향후": (
            "❍ 향후 3개월 내 중심 격자 주변 3x3 권역으로의 신규 확산위험을 "
            "지속적으로 관찰할 필요가 있음"
        ),
        "❍ (시뮬레이션 종합 의견)": (
            f"❍ (시뮬레이션 종합 의견) 중심 격자 {record.center_grid_id}를 "
            "포함한 3x3 권역에서 위험 신호가 확인되므로 중심 격자와 인접 "
            f"격자를 묶어 {metrics['priority_grade']} 대상으로 관리하고 현장 "
            "확인 결과를 격자 단위로 기록할 필요가 있다."
        ),
        "❍ 1단계:": "❍ 1단계: 3x3 권역 우선 예찰",
        "― (실행 조건)": (
            f"― (실행 조건) 종합 위험도 {metrics['risk_grade']} 또는 감염압력 "
            f"{metrics['infection_pressure']:.1f}점 이상이며 현장 확인이 "
            "완료되지 않은 경우"
        ),
        "― (실행 계획)": (
            f"― (실행 계획) 중심 격자 {record.center_grid_id} 및 주변 "
            f"{metrics['block_count'] - 1}개 격자를 대상으로 현장 예찰 또는 "
            "드론 예찰 시행"
        ),
        "❍ 2단계:": "❍ 2단계: 현장 확인 및 검경 연계",
        "― (대상 조직·알림 내용)": (
            f"― (대상 조직·알림 내용) 관할 산림부서에 {region} 3x3 예찰 "
            "검토권역의 위험도, 감염 발생 이력 및 예찰 우선순위를 공유"
        ),
        "― (현장 확인 계획)": (
            "― (현장 확인 계획) 변색목·고사목 유무를 확인하고 의심목 발견 시 "
            "시료 채취 및 전문기관 검경 의뢰"
        ),
        "❍ 3단계:": "❍ 3단계: 방제 검토 및 사후 모니터링",
        "― (방제·수종전환 검토 사항)": (
            "― (방제·수종전환 검토 사항) 현장 확인 결과와 주변 소나무류 "
            "분포를 확인한 뒤 방제 대상 지정 및 임분 관리 필요성을 검토"
        ),
        "― (행정 연계 및 후속 계획)": (
            "― (행정 연계 및 후속 계획) 예찰 결과를 격자 단위 이력으로 "
            "등록하고 인접 권역 모니터링 주기를 조정"
        ),
    }
    for prefix, value in paragraphs.items():
        set_paragraph(document, prefix, value)

    # Vercel의 DOCX→PDF 변환 환경에는 일부 원문 기호 글리프가 없어
    # ❍, ―, Ⅳ 등이 네모(속칭 엑박)로 출력될 수 있다. 문단 탐색과 내용
    # 치환을 먼저 끝낸 뒤, PDF 호환성이 높은 ASCII 표기로 정규화한다.
    replace_everywhere(
        document,
        {
            "Ⅰ": "I",
            "Ⅱ": "II",
            "Ⅲ": "III",
            "Ⅳ": "IV",
            "Ⅴ": "V",
            "❍": "-",
            "―": "-",
        },
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    report_appendix, aerial_plan = create_filled_appendix_images(
        template_path,
        appendix_directory,
        record,
        metrics,
    )
    replace_docx_media_images(
        output_path,
        {
            "word/media/image1.png": map_path,
            "word/media/image2.png": report_appendix,
            "word/media/image3.png": aerial_plan,
        },
    )


def build_prediction_render_payload(
    *,
    record: ReportRecord,
    metrics: dict[str, Any],
    source: ReportSourceData,
    map_path: Path,
    neighbor_metrics: list[dict[str, Any]] | None = None,
    start_date: str | None = None,
    end_date: str | None = None,
    title: str | None = None,
) -> dict[str, Any]:
    center_point = source.static.get("center_point_4326") or {}
    coordinates = center_point.get("coordinates") or [None, None]
    longitude = coordinates[0] if len(coordinates) > 0 else None
    latitude = coordinates[1] if len(coordinates) > 1 else None
    block_grid_ids = [int(value) for value in metrics["block_grid_ids"]]

    # 방위·거리·위험도·우선순위가 실린 인접 격자 정보가 넘어오면 그대로 쓴다.
    # 없을 때만 3x3 블록에서 격자 ID만 뽑아 쓰던 기존 동작으로 되돌아간다
    # (이 경우 표의 수치 칸은 "-"로 표시된다).
    if neighbor_metrics:
        neighbor_grids = [dict(item) for item in neighbor_metrics][:4]
    else:
        neighbor_grids = [
            {"grid_id": grid_id}
            for grid_id in block_grid_ids
            if grid_id != record.center_grid_id
        ][:4]
    end_day = 10 + (record.report_no % 18)
    center_grid = {
        "grid_id": record.center_grid_id,
        "sido_name": record.sido_name,
        "sigungu_name": record.sigungu_name,
        "risk_score": metrics["risk_score"],
        "risk_grade": metrics["risk_grade"],
        "priority_score": metrics["priority_score"],
        "priority_grade": metrics["priority_grade"],
        "infection_pressure": metrics["infection_pressure"],
        "access_score": metrics["access_score"],
        "road_distance": metrics["road_distance"],
        "road_type": metrics["road_type"],
        "environment_flag": metrics["environment_warning"],
        # report_render._percent()는 0~1 입력을 100배 하므로, 기존 DOCX가
        # 백분율 값으로 표시하던 pine_mean을 비율 입력으로 맞춘다.
        "pine_ratio": metrics["pine_mean"] / 100.0,
        "latitude": latitude,
        "longitude": longitude,
        "center_annual_count": record.annual_count,
        "center_cumulative_count": record.cumulative_count,
        "block_grid_ids": block_grid_ids,
    }
    prediction_data = {
        **center_grid,
        "center_grid_id": record.center_grid_id,
        "document_no": record.report_no,
        "year": record.year,
        "block_count": metrics["block_count"],
    }
    return {
        "report_type": "prediction",
        "document_no": record.report_no,
        "report_no": record.report_no,
        "year": record.year,
        "title": title
        or (
            f"{record.year}년 {record.sigungu_name} "
            "신규 확산위험 분석 보고서"
        ),
        # 사용자가 입력한 기간을 그대로 쓴다.
        # 값이 없을 때만 예전처럼 연도 기준으로 만들어 채운다.
        "start_date": start_date or f"{record.year}-01-01",
        "end_date": end_date or f"{record.year}-12-{end_day:02d}",
        "sido_name": record.sido_name,
        "sigungu_name": record.sigungu_name,
        "center_grid_id": record.center_grid_id,
        "center_grid_ids": [str(record.center_grid_id)],
        "block_grid_ids": block_grid_ids,
        "map_path": str(map_path),
        "prediction_data": prediction_data,
        "data": prediction_data,
        "data_summary": {
            "selected_grid_count": 1,
            "region_candidate_count": 1,
            "grid_ids": [str(record.center_grid_id)],
            "center_grid": center_grid,
            "neighbor_grids": neighbor_grids,
        },
    }


def generate_single_prediction_report(
    *,
    center_grid_id: int,
    year: int,
    output_root: Path | None = None,
    report_no: int = 1,
    zoom: int | None = None,
    candidate_metrics: dict[str, Any] | None = None,
    neighbor_metrics: list[dict[str, Any]] | None = None,
    start_date: str | None = None,
    end_date: str | None = None,
    title: str | None = None,
    client: Any | None = None,
) -> dict[str, Any]:
    if year < 2016 or year > 2100:
        raise ValueError("보고서 연도는 2016~2100 범위여야 합니다.")

    load_dotenv(BACKEND_ROOT / ".env")
    api_key = os.getenv("VWORLD_API_KEY", "").strip()
    domain = os.getenv("VWORLD_API_DOMAIN", "").strip()
    basemap = os.getenv("VWORLD_BASEMAP", "GRAPHIC").strip() or "GRAPHIC"
    validate_prediction_template(TEMPLATE_PATH)
    center_grid_id = int(center_grid_id)
    selected_zoom = zoom or int(
        os.getenv("VWORLD_ZOOM", str(DEFAULT_ZOOM))
    )
    source = load_report_source_data(
        center_grid_id,
        year,
        client=client,
    )
    historical_year = 2016 <= year <= HISTORY_LAST_YEAR
    record = ReportRecord(
        report_no=int(report_no),
        year=int(year),
        center_grid_id=center_grid_id,
        annual_count=(
            int(source.stats["center_annual_count"])
            if historical_year
            else 0
        ),
        cumulative_count=int(source.stats["center_cumulative_count"]),
        sido_name=str(source.static["sido_name"]),
        sigungu_name=str(source.static["sigungu_name"]),
    )
    metrics = apply_candidate_metrics(
        calculate_metrics(record, source),
        candidate_metrics,
    )
    root = (
        output_root.resolve()
        if output_root is not None
        else DEFAULT_OUTPUT_ROOT.resolve()
    )
    docx_directory = root / "docx"
    pdf_directory = root / "pdf"
    map_directory = root / "maps"
    appendix_directory = root / "appendices"
    for directory in [
        root,
        docx_directory,
        pdf_directory,
        map_directory,
        appendix_directory,
    ]:
        directory.mkdir(parents=True, exist_ok=True)

    base_name = sanitize_filename(
        f"{record.report_no:02d}_{record.year}_"
        "소나무재선충병_신규확산위험분석보고서_"
        f"{record.sido_name}_{record.sigungu_name}_"
        f"격자{record.center_grid_id}"
    )
    map_path = map_directory / f"{base_name}.png"
    docx_path = docx_directory / f"{base_name}.docx"
    pdf_path = pdf_directory / f"{base_name}.pdf"
    build_vworld_overlay_map(
        output_path=map_path,
        record=record,
        source=source,
        api_key=api_key,
        domain=domain,
        zoom=selected_zoom,
        basemap=basemap,
    )
    create_docx(
        template_path=TEMPLATE_PATH,
        output_path=docx_path,
        map_path=map_path,
        appendix_directory=appendix_directory,
        record=record,
        metrics=metrics,
        start_date=start_date,
        end_date=end_date,
    )
    render_payload = build_prediction_render_payload(
        record=record,
        metrics=metrics,
        source=source,
        map_path=map_path,
        neighbor_metrics=neighbor_metrics,
        start_date=start_date,
        end_date=end_date,
        title=title,
    )
    from app.services.report_render.renderer import render_report_pdf

    pdf_path.write_bytes(render_report_pdf("prediction", render_payload))

    for label, path in {
        "지도": map_path,
        "DOCX": docx_path,
        "PDF": pdf_path,
    }.items():
        if not path.is_file():
            raise RuntimeError(f"{label} 생성 결과를 찾을 수 없습니다: {path}")

    return {
        "center_grid_id": record.center_grid_id,
        "year": record.year,
        "sido_name": record.sido_name,
        "sigungu_name": record.sigungu_name,
        "annual_count": record.annual_count,
        "cumulative_count": record.cumulative_count,
        "risk_score": round(float(metrics["risk_score"]), 2),
        "risk_grade": metrics["risk_grade"],
        "priority_score": round(float(metrics["priority_score"]), 2),
        "priority_grade": metrics["priority_grade"],
        "block_grid_ids": metrics["block_grid_ids"],
        "map_path": str(map_path),
        "docx_path": str(docx_path),
        "pdf_path": str(pdf_path),
    }
