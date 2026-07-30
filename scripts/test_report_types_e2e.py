#!/usr/bin/env python3
"""세 보고서 유형의 초안 저장부터 산출물 생성까지 검증한다."""
from __future__ import annotations

import json
import os
import sys
import zipfile
from pathlib import Path

import fitz
from dotenv import load_dotenv
from fastapi.testclient import TestClient
from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
BACKEND_ROOT = PROJECT_ROOT / "rag-backend"
RESULT_PATH = (
    PROJECT_ROOT
    / "data"
    / "precompute_logs"
    / "report_types_e2e_20260726.json"
)
sys.path.insert(0, str(BACKEND_ROOT))

load_dotenv(BACKEND_ROOT / ".env")
load_dotenv(PROJECT_ROOT / ".env")
if not os.getenv("VWORLD_API_KEY"):
    os.environ["VWORLD_API_KEY"] = os.getenv("VITE_VWORLD_API_KEY", "")
if not os.getenv("VWORLD_API_DOMAIN"):
    os.environ["VWORLD_API_DOMAIN"] = "localhost"

from app.api.report_drafts import _apply_template  # noqa: E402
from app.api.auth import get_current_user  # noqa: E402
from app.main import app  # noqa: E402
from app.models.user import User  # noqa: E402
from app.services.report_draft_service import (  # noqa: E402
    build_xlsx,
    create_draft,
    load_draft,
)


def validate_docx(path: Path) -> bool:
    with zipfile.ZipFile(path) as archive:
        return "word/document.xml" in archive.namelist()


def main() -> int:
    test_user = User(
        id=0,
        email="codex-e2e@local.test",
        password_hash="not-used",
        name="Codex E2E",
        organization="local-test",
        role="manager",
        is_active=True,
    )
    app.dependency_overrides[get_current_user] = lambda: test_user
    client = TestClient(app)
    results: list[dict[str, object]] = []
    try:
        for report_type in ("prediction", "field_survey", "control"):
            payload = {
                "report_type": report_type,
                "title": f"E2E {report_type} 20260726",
                "year": 2021,
                "start_date": "2021-12-01",
                "end_date": "2021-12-31",
                "sido_name": "경상북도",
                "sigungu_name": "포항시 북구",
                "center_grid_ids": ["1001377"],
                "include_sections": {
                    "risk_summary": True,
                    "priority_summary": True,
                    "infection_history": True,
                    "workforce_plan": False,
                    "control_scenario": False,
                },
                "user_notes": "3개 report_type E2E 검증",
            }
            draft = create_draft(payload, created_by="codex-e2e@local.test")
            output = _apply_template(report_type, draft["draft_id"])
            stored = load_draft(draft["draft_id"])
            xlsx_path = build_xlsx(stored)
            pdf_path = Path(output["pdf_path"])
            docx_path = Path(output["docx_path"])

            with fitz.open(pdf_path) as pdf:
                page_count = len(pdf)
            workbook = load_workbook(xlsx_path, read_only=True)
            sheet_names = workbook.sheetnames
            workbook.close()

            draft_id = draft["draft_id"]
            responses = {
                "get": client.get(f"/api/report-drafts/{draft_id}"),
                "preview_pdf": client.get(
                    f"/api/report-drafts/{draft_id}/preview/pdf"
                ),
                **{
                    f"export_{file_format}": client.post(
                        f"/api/report-drafts/{draft_id}/export/{file_format}"
                    )
                    for file_format in ("pdf", "docx", "xlsx")
                },
            }
            results.append(
                {
                    "report_type": report_type,
                    "draft_id": draft_id,
                    "stored_template_output": (
                        stored.get("template_output") == output
                    ),
                    "pdf": {
                        "exists": pdf_path.is_file(),
                        "bytes": pdf_path.stat().st_size,
                        "pages": page_count,
                    },
                    "docx": {
                        "exists": docx_path.is_file(),
                        "bytes": docx_path.stat().st_size,
                        "valid_package": validate_docx(docx_path),
                    },
                    "xlsx": {
                        "exists": xlsx_path.is_file(),
                        "bytes": xlsx_path.stat().st_size,
                        "sheets": sheet_names,
                    },
                    "http": {
                        name: {
                            "status": response.status_code,
                            "content_type": response.headers.get(
                                "content-type", ""
                            ),
                            "bytes": len(response.content),
                        }
                        for name, response in responses.items()
                    },
                    "return_keys": sorted(output),
                    "field_survey_linked": output.get(
                        "field_survey_linked"
                    ),
                },
            )
    finally:
        app.dependency_overrides.clear()

    RESULT_PATH.parent.mkdir(parents=True, exist_ok=True)
    RESULT_PATH.write_text(
        json.dumps(results, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(json.dumps(results, ensure_ascii=False, indent=2))

    valid = all(
        item["stored_template_output"]
        and item["pdf"]["exists"]
        and item["pdf"]["pages"] > 0
        and item["docx"]["exists"]
        and item["docx"]["valid_package"]
        and item["xlsx"]["exists"]
        and all(
            response["status"] == 200
            for response in item["http"].values()
        )
        for item in results
    )
    return 0 if valid else 1


if __name__ == "__main__":
    raise SystemExit(main())
