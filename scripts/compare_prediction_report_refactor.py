#!/usr/bin/env python3
"""기존 GeoPandas 보고서 경로와 신규 Supabase 경로를 로컬에서 비교한다."""

from __future__ import annotations

import importlib.util
import argparse
import json
import math
import sys
import zipfile
from pathlib import Path
from typing import Any

from PIL import Image, ImageChops
from pyproj import Transformer


PROJECT_ROOT = Path(__file__).resolve().parents[1]
BACKEND_ROOT = PROJECT_ROOT / "rag-backend"
OUTPUT_BASE = PROJECT_ROOT / "data" / "precompute_logs" / "report_refactor_compare"
ZOOM = 10

sys.path.insert(0, str(BACKEND_ROOT))

from app.services import prediction_report_generator as new  # noqa: E402


def load_module(name: str, path: Path) -> Any:
    spec = importlib.util.spec_from_file_location(name, path)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"모듈을 불러올 수 없습니다: {path}")
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module
    spec.loader.exec_module(module)
    return module


def close(a: float, b: float, tolerance: float = 1e-9) -> bool:
    return math.isclose(float(a), float(b), rel_tol=0.0, abs_tol=tolerance)


def docx_text(docx_path: Path) -> list[str]:
    from docx import Document

    document = Document(docx_path)
    values: list[str] = []
    values.extend(paragraph.text for paragraph in document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                values.extend(paragraph.text for paragraph in cell.paragraphs)
    for section in document.sections:
        values.extend(paragraph.text for paragraph in section.header.paragraphs)
        values.extend(paragraph.text for paragraph in section.footer.paragraphs)
    return values


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--grid-id", type=int, default=321342)
    parser.add_argument("--year", type=int, default=2021)
    arguments = parser.parse_args()
    center_grid_id = arguments.grid_id
    year = arguments.year
    output_root = OUTPUT_BASE / f"{center_grid_id}_{year}"
    output_root.mkdir(parents=True, exist_ok=True)
    old_single = load_module(
        "old_prediction_report_single",
        PROJECT_ROOT / "scripts" / "prediction_report_single.py",
    )
    old_batch = old_single._load_batch_module()

    terrain_by_id, terrain_by_corner = old_batch.load_terrain_index(
        old_single.TERRAIN_PATH
    )
    candidate_row, _ = old_single._load_candidate_row(center_grid_id)
    infection = old_batch.load_infection_history(old_single.INFECTION_PATH)
    infection_by_id = infection.set_index("id", drop=False)
    cells = old_batch.get_3x3_cells(
        center_grid_id,
        terrain_by_id,
        terrain_by_corner,
    )
    center_history = (
        infection_by_id.loc[center_grid_id]
        if center_grid_id in infection_by_id.index
        else None
    )
    if center_history is not None and hasattr(center_history, "iloc"):
        if getattr(center_history, "ndim", 1) > 1:
            center_history = center_history.iloc[0]
    old_record = old_batch.ReportRecord(
        report_no=1,
        year=year,
        center_grid_id=center_grid_id,
        annual_count=(
            int(center_history.get(f"infection_count_{year}", 0) or 0)
            if center_history is not None
            else 0
        ),
        cumulative_count=(
            int(center_history.get("infection_count_2016_2021", 0) or 0)
            if center_history is not None
            else 0
        ),
        sido_name=str(candidate_row["sido_name"]),
        sigungu_name=str(candidate_row["sigungu_name"]),
    )
    old_metrics = old_batch.calculate_metrics(
        old_record,
        cells,
        infection_by_id,
    )
    old_metrics = old_single._apply_candidate_metrics(
        old_batch,
        old_metrics,
        candidate_row,
    )

    source = new.load_report_source_data(center_grid_id, year)
    new_record = new.ReportRecord(
        report_no=1,
        year=year,
        center_grid_id=center_grid_id,
        annual_count=int(source.stats["center_annual_count"]),
        cumulative_count=int(source.stats["center_cumulative_count"]),
        sido_name=str(source.static["sido_name"]),
        sigungu_name=str(source.static["sigungu_name"]),
    )
    candidate_metrics = {
        "risk_score": candidate_row.get("risk_score"),
        "risk_grade": candidate_row.get("risk_grade"),
        "access_score_v3": candidate_row.get("access_score_v3"),
    }
    new_metrics = new.apply_candidate_metrics(
        new.calculate_metrics(new_record, source),
        candidate_metrics,
    )

    background = Image.new(
        "RGBA",
        (new.MAP_WIDTH, new.MAP_HEIGHT),
        (232, 238, 228, 255),
    )
    old_map = output_root / "old_overlay.png"
    new_map = output_root / "new_overlay.png"
    original_request = old_batch.request_vworld_map
    old_batch.request_vworld_map = lambda **_: background.copy()
    try:
        old_batch.build_vworld_overlay_map(
            output_path=old_map,
            record=old_record,
            cells=cells,
            infection=infection,
            api_key="fixed",
            domain="fixed",
            zoom=ZOOM,
            basemap="GRAPHIC",
        )
    finally:
        old_batch.request_vworld_map = original_request
    _, new_trace = new.render_vworld_overlay(
        background,
        new_record,
        source,
        ZOOM,
    )
    new_image, _ = new.render_vworld_overlay(
        background,
        new_record,
        source,
        ZOOM,
    )
    new_image.save(new_map)

    transformer = Transformer.from_crs(
        "EPSG:5186",
        "EPSG:4326",
        always_xy=True,
    )
    center_cell = next(
        cell for cell in cells if cell.grid_id == center_grid_id
    )
    center_point = center_cell.geometry.centroid
    center_lon, center_lat = transformer.transform(
        center_point.x,
        center_point.y,
    )
    old_block_trace: dict[str, list[list[tuple[float, float]]]] = {}
    for cell in cells:
        old_block_trace[str(cell.grid_id)] = [
            old_batch.geometry_to_image_points(
                cell.geometry,
                center_lon,
                center_lat,
                ZOOM,
                new.MAP_WIDTH,
                new.MAP_HEIGHT,
                transformer,
            )
        ]
    coordinate_deltas = []
    for grid_id, old_rings in old_block_trace.items():
        new_rings = new_trace["block_polygons"][grid_id]
        for old_ring, new_ring in zip(old_rings, new_rings):
            for old_point, new_point in zip(old_ring, new_ring):
                coordinate_deltas.extend(
                    [
                        abs(old_point[0] - new_point[0]),
                        abs(old_point[1] - new_point[1]),
                    ]
                )

    old_docx = output_root / "old.docx"
    new_docx = output_root / "new.docx"
    old_batch.create_docx(
        template_path=old_single.TEMPLATE_PATH,
        output_path=old_docx,
        map_path=old_map,
        appendix_dir=output_root / "old_appendices",
        record=old_record,
        cells=cells,
        metrics=old_metrics,
    )
    new.create_docx(
        template_path=new.TEMPLATE_PATH,
        output_path=new_docx,
        map_path=new_map,
        appendix_directory=output_root / "new_appendices",
        record=new_record,
        metrics=new_metrics,
    )

    difference = ImageChops.difference(
        Image.open(old_map).convert("RGB"),
        Image.open(new_map).convert("RGB"),
    )
    changed_pixels = sum(
        1 for pixel in difference.getdata() if pixel != (0, 0, 0)
    )
    metric_keys = [
        "block_count",
        "block_annual",
        "block_cumulative",
        "active_count",
        "pine_mean",
        "elevation_mean",
        "slope_mean",
        "infection_pressure",
        "access_score",
        "road_distance",
        "risk_score",
        "risk_stage",
        "risk_grade",
        "priority_score",
        "priority_grade",
    ]
    metric_results = {}
    for key in metric_keys:
        old_value = old_metrics[key]
        new_value = new_metrics[key]
        equal = (
            close(old_value, new_value)
            if isinstance(old_value, (int, float))
            and isinstance(new_value, (int, float))
            else old_value == new_value
        )
        metric_results[key] = {
            "old": old_value,
            "new": new_value,
            "equal": equal,
        }

    old_text = docx_text(old_docx)
    new_text = docx_text(new_docx)
    result = {
        "grid_id": center_grid_id,
        "year": year,
        "record": {
            "annual_count": [
                old_record.annual_count,
                new_record.annual_count,
            ],
            "cumulative_count": [
                old_record.cumulative_count,
                new_record.cumulative_count,
            ],
        },
        "metrics": metric_results,
        "block_grid_ids_equal": (
            old_metrics["block_grid_ids"] == new_metrics["block_grid_ids"]
        ),
        "docx_text_equal": old_text == new_text,
        "docx_text_count": [len(old_text), len(new_text)],
        "remaining_placeholders": {
            "old": sum("[" in value and "]" in value for value in old_text),
            "new": sum("[" in value and "]" in value for value in new_text),
        },
        "overlay": {
            "changed_pixels": changed_pixels,
            "total_pixels": new.MAP_WIDTH * new.MAP_HEIGHT,
            "pixel_equal": changed_pixels == 0,
            "block_coordinate_max_delta_px": max(
                coordinate_deltas,
                default=0.0,
            ),
        },
    }
    output_json = output_root / "comparison.json"
    output_json.write_text(
        json.dumps(result, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(json.dumps(result, ensure_ascii=False, indent=2))

    success = (
        all(item["equal"] for item in metric_results.values())
        and result["block_grid_ids_equal"]
        and result["docx_text_equal"]
        and result["overlay"]["pixel_equal"]
    )
    return 0 if success else 1


if __name__ == "__main__":
    raise SystemExit(main())
